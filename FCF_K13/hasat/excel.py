"""hasat.pkl → HASAT_TARAMA.xlsx"""
import os, re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule

KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__))); os.chdir(KOK)
S = pd.read_pickle('hasat/hasat.pkl')
GET = [c for c in S.columns if str(c).startswith('Getiri')]
SIRA = {'GÜÇLÜ ADAY': 0, 'ADAY': 1, 'İZLE': 2, 'UYMUYOR': 3, 'TUZAK ŞÜPHESİ': 4, 'ELENDİ': 5}
S['_s'] = S.Kategori.map(SIRA)
S = S.sort_values(['_s', 'HASAT PUANI'], ascending=[True, False])

ORT = ['Kod', 'Ünvan', 'Sektör', 'Tip', 'Baz', 'Kategori', 'HASAT PUANI', 'Yatırım döngüsü', 'CAPEX zirve yılı', 'Borç durumu',
       'Hasat getirisi', 'Borç sonrası hasat getirisi', 'FCF getirisi', 'P/FCF', 'PD/DD', 'FD/FAVÖK', 'Temettü verimi T', 'Temettü ödenen yıl',
       'Satış değişimi 23→T', 'Brüt kâr değişimi 23→T', 'Çöküş uyarısı', 'PD'] + GET
PUAN_S = ['Döngü (25)', 'Nakit (15)', 'Amortisman kalkanı (10)', 'Borç (15)', 'Brüt kâr reel (10)', 'Temettü (10)', 'Ucuzluk (15)']
PUAN_G = ['Döngü (20)', 'Nakit (15)', 'Borç (20)', 'Gelir reel (10)', 'Temettü (10)', 'Ucuzluk PD/DD (25)']
SERI = []
for y in ['23', '24', '25', '26']:
    SERI += [f'CFO {y}', f'CAPEX {y}', f'Amort {y}', f'Kira {y}', f'FCF {y}']
SERI += ['Bakım CAPEX (vekil)', 'Hasat nakdi T', 'Amort/CAPEX T', 'FCF/Net kâr T'] + \
        [f'CAPEX/Amort {y}' for y in ['23', '24', '25', '26']] + [f'CAPEX/Satış {y}' for y in ['23', '24', '25', '26']] + \
        [f'Satış {y}' for y in ['23', '24', '25', '26']] + [f'Brüt {y}' for y in ['23', '24', '25', '26']] + \
        [f'FAVÖK {y}' for y in ['23', '24', '25', '26']] + [f'Net kâr {y}' for y in ['23', '24', '25', '26']] + \
        [f'Net borç {y}' for y in ['23', '24', '25', '26']] + ['Net borç/FAVÖK 23', 'Net borç/FAVÖK T', 'Ödenen faiz 26', 'Ödenen faiz/CFO T'] + \
        [f'Ödenen tem {y}' for y in ['23', '24', '25', '26']] + ['Satış 23 USD', 'Satış 26 USD', 'Brüt 23 USD', 'Brüt 26 USD',
        'Net borç 23 USD', 'Net borç 26 USD', 'Özkaynak (ana ort.)', 'FCF statü 25', 'FCF statü 26', 'Eleme']
SERI_G = [c.replace('CAPEX/Amort', 'CAPEX/CFO') for c in SERI] + ['PD/Satış (kira çarpanı vekili)']

s = S[S.Grup == 'Şirket']; g = S[S.Grup == 'GYO']
aday = s[s.Kategori.isin(['GÜÇLÜ ADAY', 'ADAY', 'İZLE'])]
kural = open('hasat/KURAL_SETI.md').read()
OK = [(l.strip('# ').strip(), '') if l.startswith('#') else ('', l) for l in kural.split('\n') if l.strip()]
ozet = pd.concat([S.groupby(['Grup', 'Kategori']).size().rename('Şirket').reset_index()])

with pd.ExcelWriter('HASAT_TARAMA.xlsx', engine='openpyxl') as w:
    pd.DataFrame(OK, columns=['Başlık', 'Kural seti']).to_excel(w, sheet_name='KURAL_SETI', index=False)
    ozet.to_excel(w, sheet_name='OZET', index=False)
    aday[ORT + PUAN_S + SERI].to_excel(w, sheet_name='ADAYLAR', index=False)
    g[ORT + ['PD/Satış (kira çarpanı vekili)'] + PUAN_G + [c for c in SERI_G if c in g.columns and c != 'PD/Satış (kira çarpanı vekili)']].to_excel(w, sheet_name='GYO', index=False)
    s[s.Kategori == 'TUZAK ŞÜPHESİ'][ORT + ['Eleme'] + SERI].to_excel(w, sheet_name='TUZAK', index=False)
    s[s.Kategori.isin(['UYMUYOR', 'ELENDİ'])][ORT + ['Eleme'] + PUAN_S + SERI].to_excel(w, sheet_name='DIGER', index=False)

wb = load_workbook('HASAT_TARAMA.xlsx')
YUZDE = re.compile(r'getirisi|verimi|değişimi|Hasat getirisi')
for ws in wb:
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78'); c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 42
    if ws.title == 'KURAL_SETI':
        ws.column_dimensions['A'].width = 40; ws.column_dimensions['B'].width = 160
        for r in ws.iter_rows(min_row=2):
            if r[0].value: r[0].font = Font(bold=True, color='1F4E78', size=12)
            r[1].alignment = Alignment(wrap_text=True, vertical='top')
        continue
    if ws.title == 'OZET':
        for col in 'ABC': ws.column_dimensions[col].width = 18
        continue
    ws.freeze_panes = 'C2'
    for col in ws.columns:
        hd = str(col[0].value); L = col[0].column_letter
        ws.column_dimensions[L].width = 32 if hd == 'Ünvan' else 26 if hd in ('Yatırım döngüsü', 'Borç durumu', 'Sektör') else 40 if hd in ('Eleme', 'Çöküş uyarısı') else 12
        fmt = None
        if YUZDE.search(hd) or hd.startswith('Getiri'): fmt = '0.0%'
        elif '/' in hd or hd.startswith('P/') or hd.startswith('PD/') or hd.startswith('FD/'): fmt = '0.00'
        elif any(hd.startswith(k) for k in ('CFO', 'CAPEX ', 'Amort ', 'Kira', 'FCF ', 'Satış ', 'Brüt ', 'FAVÖK', 'Net ', 'Ödenen', 'Bakım', 'Hasat nakdi', 'PD', 'Özkaynak')):
            fmt = '#,##0.0;[Red]-#,##0.0'
        if fmt and not hd.startswith('FCF statü'):
            for c in col[1:]:
                if isinstance(c.value, (int, float)): c.number_format = fmt
    ws.auto_filter.ref = ws.dimensions
    kat = next((c.column_letter for c in ws[1] if c.value == 'Kategori'), None)
    if kat:
        for t, clr in (('GÜÇLÜ ADAY', 'C6EFCE'), ('ADAY', 'DDEBF7'), ('İZLE', 'FFF2CC'), ('TUZAK ŞÜPHESİ', 'F8CBAD')):
            ws.conditional_formatting.add(f'{kat}2:{kat}{ws.max_row}', CellIsRule(operator='equal', formula=[f'"{t}"'], fill=PatternFill('solid', fgColor=clr)))
wb.save('HASAT_TARAMA.xlsx')
print(ozet.to_string())
