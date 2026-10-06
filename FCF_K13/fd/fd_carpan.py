"""FD/FAVÖK ve FD/Faaliyet kârı (temiz EFK) — tüm evren.
FD = güncel piyasa değeri (PFCF_GETIRI, 06.10.2026) + net borç (2026/06 bilanço; net nakit FD'yi düşürür).
FAVÖK = Funnel standardı FAVÖK_STD TTM 2026/06 = brüt kâr − pazarlama − genel yönetim − Ar-Ge + amortisman.
Faaliyet kârı (temiz EFK) = FAVÖK_STD − amortisman = brüt kâr − pazarlama − genel yönetim − Ar-Ge (TTM).
  Hariç: net parasal pozisyon kazancı (TMS 29), finansman gelir/gideri (faiz, kur farkı), yatırım faaliyetlerinden gelirler
  (fon/menkul kıymet getirisi), esas faaliyetlerden diğer gelir/giderler, özkaynak yöntemi payları.
Kaynak: FCF_MASTER FUNNEL_GIRDI (Evo, 2026/06 TTM). Payda ≤ 0 → anlamsız."""
import os
import pandas as pd, numpy as np
os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
G = pd.read_excel('FCF_MASTER.xlsx', sheet_name='FUNNEL_GIRDI'); G = G[G['Dönem'] == '2026/06'].set_index('Kod')
A = pd.read_excel('FCF_MASTER.xlsx', sheet_name='ANA', header=None).iloc[4:, :5]; A.columns = ['Kod', 'Ünvan', 'Kapsam', 'Tip', 'Sektör']; A = A.set_index('Kod')
F = pd.read_excel('FCF_MASTER.xlsx', sheet_name='FUNNEL'); F = F[F['Dönem'] == '2026/06'].set_index('Kod')
P = pd.read_excel('PFCF_GETIRI.xlsx', sheet_name='P_FCF', header=None)
h = next(i for i in range(10) if 'Kod' in [str(x) for x in P.iloc[i]]); P.columns = P.iloc[h]; P = P.iloc[h + 1:]; P = P[P.Kod.notna()].set_index('Kod')
mn = lambda s: pd.to_numeric(s, errors='coerce') / 1e6
D = pd.DataFrame(index=G.index)
D['Ünvan'] = A['Ünvan']; D['Sektör'] = A['Sektör']; D['Kapsam'] = A['Kapsam']; D['Şablon'] = G['finansal_tablo_sablonu']
D['PD (güncel)'] = pd.to_numeric(P['Piyasa değeri (mn TL)'], errors='coerce').reindex(D.index)
D['PD'] = D['PD (güncel)'].fillna(mn(G.piyasa_degeri))
D['Net borç'] = mn(G.net_borc_cur)
D['FD'] = D.PD + D['Net borç'].fillna(0)
D['Satış TTM'] = mn(G.satis_cur); D['Brüt kâr TTM'] = mn(G.brut_cur); D['Faaliyet giderleri TTM'] = mn(G.opex_raw_cur)
D['Amortisman TTM'] = mn(G.amort_cur); D['FAVÖK TTM'] = mn(G.favok_cur)
D['Faaliyet kârı (temiz EFK) TTM'] = D['Brüt kâr TTM'] + D['Faaliyet giderleri TTM']
D['Net kâr TTM (raporlanan)'] = mn(G.net_total_cur)
def oran(a, b): return np.where((b > 0) & a.notna(), a / b, np.nan)
D['FD/FAVÖK'] = oran(D.FD, D['FAVÖK TTM'])
D['FD/Faaliyet kârı'] = oran(D.FD, D['Faaliyet kârı (temiz EFK) TTM'])
D['F/K (raporlanan)'] = oran(D.PD, D['Net kâr TTM (raporlanan)'])
D['Not'] = np.select([D['FAVÖK TTM'] <= 0, D['Faaliyet kârı (temiz EFK) TTM'] <= 0, D.FD <= 0],
                     ['FAVÖK negatif → çarpan anlamsız', 'Faaliyet kârı negatif → FD/Faaliyet kârı anlamsız', 'FD negatif (net nakit > PD)'], '')
D.loc[D.FD <= 0, ['FD/FAVÖK', 'FD/Faaliyet kârı']] = np.nan
D['Funnel MIN (2026/06)'] = F['Funnel MIN']
D = D[D.Şablon == 'default']
DIGER = D[D.Kapsam != 'STANDART']      # Tier 0 (GYO vb.) ve Tier 0,5 (aracı kurum/yatırım holdingi): FD/FAVÖK ekonomik anlam taşımaz
D = D[D.Kapsam == 'STANDART']
D = D.sort_values('FD/Faaliyet kârı')
D['Sektör medyanı FD/Faaliyet kârı'] = D.groupby('Sektör')['FD/Faaliyet kârı'].transform('median')
D['Sektör medyanı FD/FAVÖK'] = D.groupby('Sektör')['FD/FAVÖK'].transform('median')
D.reset_index().to_excel('FD_CARPANLARI.xlsx', sheet_name='FD_CARPAN', index=False)
SEK = D.groupby('Sektör').agg(Şirket=('FD', 'size'), **{'Medyan FD/FAVÖK': ('FD/FAVÖK', 'median'), 'Medyan FD/Faaliyet kârı': ('FD/Faaliyet kârı', 'median')}).sort_values('Medyan FD/Faaliyet kârı')
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
with pd.ExcelWriter('FD_CARPANLARI.xlsx', engine='openpyxl', mode='a') as w:
    SEK.reset_index().to_excel(w, sheet_name='SEKTOR', index=False)
    DIGER.reset_index().to_excel(w, sheet_name='KAPSAM_DISI', index=False)
    pd.DataFrame({'Tanım': __doc__.split('\n')}).to_excel(w, sheet_name='OKUBENI', index=False)
wb = load_workbook('FD_CARPANLARI.xlsx')
for ws in wb:
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78'); c.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = 'B2'
    for col in ws.columns:
        hd = str(col[0].value); ws.column_dimensions[col[0].column_letter].width = 32 if hd in ('Ünvan', 'Not', 'Tanım') else 14
        if ws.title == 'OKUBENI': ws.column_dimensions['A'].width = 150
        for c in col[1:]:
            if isinstance(c.value, float): c.number_format = '0.0%' if 'payı' in hd else '0.0' if ('/' in hd or 'medyan' in hd.lower()) else '#,##0'
    ws.auto_filter.ref = ws.dimensions
wb.save('FD_CARPANLARI.xlsx')
v = D[D['FD/Faaliyet kârı'].notna()]
print(len(D), 'şirket;', D['FD/FAVÖK'].notna().sum(), 'FD/FAVÖK anlamlı;', len(v), 'FD/Faaliyet kârı anlamlı')
print('Medyan FD/FAVÖK', round(D['FD/FAVÖK'].median(), 1), '| Medyan FD/Faaliyet kârı', round(v['FD/Faaliyet kârı'].median(), 1))
print(v[['FD/FAVÖK', 'FD/Faaliyet kârı', 'F/K (raporlanan)', 'Funnel MIN (2026/06)', 'Sektör']].head(25).round(1).to_string())
print(SEK.round(1).to_string())
print('FD negatif:', (D.FD <= 0).sum(), '| FAVÖK negatif:', (D['FAVÖK TTM'] <= 0).sum(), '| Faaliyet kârı negatif:', (D['Faaliyet kârı (temiz EFK) TTM'] <= 0).sum())
