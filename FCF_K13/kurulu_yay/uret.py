"""Kurulu Yay — nihai: 4 geçişin tamamıyla eğit, 2026/06 verisiyle bir sonraki bilanço (2026/09) için tara → KURULU_YAY.xlsx"""
import os, sys
import numpy as np, pandas as pd
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ky
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule

L, D = ky.L, ky.D
M = ky.egit(L)
SON = D[-1]
X = ky.tahmin(M, L[L.D == SON])
S = ky.sirket(X, n_mc=8000)
F = pd.read_excel('FCF_MASTER.xlsx', sheet_name='FUNNEL')
F = F[F['Dönem'] == SON].set_index('Kod')
A = pd.read_excel('FCF_MASTER.xlsx', sheet_name='ANA', header=None).iloc[4:, :5]; A.columns = ['Kod', 'Ünvan', 'Kapsam', 'Tip', 'Sektör']
A = A.set_index('Kod')
FC = pd.read_excel('FCF_MASTER.xlsx', sheet_name='FCF', keep_default_na=False)
fc = FC[FC['Dönem'] == SON].set_index('Kod')
S['Ünvan'] = S.Kod.map(A['Ünvan']); S['Sektör'] = S.Kod.map(A['Sektör'])
S['Funnel MIN (2026/06)'] = S.Kod.map(F['Funnel MIN']); S['Funnel MAX'] = S.Kod.map(F['Funnel MAX']); S['Kategori (2026/06)'] = S.Kod.map(F['Kategori'])
S['Beklenen Funnel'] = S['Funnel MIN (2026/06)'] + S.E_delta
S['FCF TTM 2026/06'] = S.Kod.map(pd.to_numeric(fc['FCF · TTM (mn TL)'], errors='coerce')); S['FCF statü'] = S.Kod.map(fc['Statü'])
S['Önceki Funnel (2026/03)'] = S.Kod.map(L[L.D == D[-2]].groupby('Kod').p.sum())
def sinif(r):
    if r.E_delta >= 4 and r.Asagi <= 2.5: return 'GÜÇLÜ YAY (dengeli)'
    if r.E_delta >= 4: return 'YAY — yüksek potansiyel, yüksek risk'
    if r.E_delta >= 2: return 'ORTA YAY'
    if r.E_delta <= -3: return 'TERS YAY (düşüş riski)'
    return ''
S['Sınıf'] = S.apply(sinif, axis=1)
S = S.sort_values('E_delta', ascending=False)
S['Sıra'] = range(1, len(S) + 1)
S['Not'] = np.where(S.M10_payi > 0.5, 'Beklenen kazancın yarıdan fazlası M10 (alacak/satış) — oynak metrik', '')
KOL = ['Sıra', 'Kod', 'Ünvan', 'Sektör', 'Sınıf', 'E_delta', 'P_sicrama', 'P_dusus', 'Yukari', 'Asagi', 'P_kat', 'Funnel MIN (2026/06)', 'Beklenen Funnel',
       'Önceki Funnel (2026/03)', 'Kategori (2026/06)', 'Surukleyici', 'Risk', 'Not', 'FCF TTM 2026/06', 'FCF statü', 'Funnel MAX']
AD = {'E_delta': 'KURULU YAY skoru = E[ΔFunnel]', 'P_sicrama': 'P(Δ ≥ +8)', 'P_dusus': 'P(Δ ≤ −8)', 'Yukari': 'Beklenen kazanç',
      'Asagi': 'Beklenen kayıp', 'P_kat': 'P(bir üst kategoriye geçiş)', 'Surukleyici': 'Yukarı sürükleyiciler: metrik puan→hedef (olasılık)', 'Risk': 'Aşağı riskler: metrik puan→düşüş (olasılık)'}
Y = S[KOL].rename(columns=AD)

# metrik detayı
X['Metrik'] = X.m; X['Değer'] = X.v; X['Puan'] = X.p
X['Üst banda uzaklık (σ)'] = X.z_up; X['Alt banda uzaklık (σ)'] = X.z_dn; X['Son çeyrek ivmesi (σ)'] = X.mom
X['P yukarı'] = X.P_up; X['Kazanç'] = X.kazanc; X['P aşağı'] = X.P_dn; X['Kayıp'] = X.g_dn; X['Rejim (taban oran)'] = X.rejim.map({True: 'EVET', False: ''})
MD = X[X.Kod.isin(S.head(60).Kod)].copy(); MD['_s'] = MD.Kod.map(S.set_index('Kod').Sıra)
MD = MD.sort_values(['_s', 'm'])[['Kod', 'Metrik', 'Değer', 'Puan', 'Üst banda uzaklık (σ)', 'Alt banda uzaklık (σ)', 'Son çeyrek ivmesi (σ)',
                                   'P yukarı', 'Kazanç', 'P aşağı', 'Kayıp', 'Rejim (taban oran)']]

BT = pd.read_csv('kurulu_yay/backtest_sonuc.csv')
P = pd.read_pickle('kurulu_yay/backtest_tahmin.pkl').dropna(subset=['Gercek_delta'])
P['Desil'] = P.groupby('Test').E_delta.transform(lambda x: pd.qcut(x.rank(method='first'), 10, labels=False) + 1)
KAL = P.groupby('Desil').agg(**{'Tahmin E[Δ]': ('E_delta', 'mean'), 'Gerçekleşen Δ': ('Gercek_delta', 'mean'),
                               'Tahmin P(≥+8)': ('P_sicrama', 'mean'), 'Gerçekleşen oran (≥+8)': ('Gercek_delta', lambda x: (x >= 8).mean()),
                               'Gerçekleşen oran (≤−8)': ('Gercek_delta', lambda x: (x <= -8).mean()), 'Gözlem': ('Kod', 'count')}).reset_index()
MOD = []
for m, w in M.items():
    MOD.append(dict(Metrik=m, **{f'yukarı_{k}': round(v, 3) for k, v in zip(['sabit', 'log(1+z_up)', 'ivme', 'ivme_var'], w['up'])} if w['up'] is not None else {},
                    **{f'aşağı_{k}': round(v, 3) for k, v in zip(['sabit', 'log(1+z_dn)', 'ivme', 'ivme_var'], w['dn'])} if w['dn'] is not None else {},
                    Rejim_yukarı_oran=round(w['rejim_up'], 3), Rejim_aşağı_oran=round(w['rejim_dn'], 3), Ölçek_σ=round(ky.SKALA[m], 4)))
MET = [l.rstrip('\n') for l in open('kurulu_yay/METODOLOJI.md')]
with pd.ExcelWriter('KURULU_YAY.xlsx', engine='openpyxl') as w:
    pd.DataFrame({'Metodoloji': MET}).to_excel(w, sheet_name='METODOLOJI', index=False)
    BT.to_excel(w, sheet_name='BACKTEST', index=False)
    KAL.to_excel(w, sheet_name='KALIBRASYON', index=False)
    KA = S[S['Funnel MIN (2026/06)'] < 70].sort_values('P_kat', ascending=False).copy()
    KA['Hedef kategori'] = np.where(KA['Funnel MIN (2026/06)'] >= 60, 'ANA LİSTE (≥70)', 'İZLEME (≥60)')
    KA['Sıra'] = range(1, len(KA) + 1)
    KA[['Sıra', 'Kod', 'Ünvan', 'Sektör', 'Funnel MIN (2026/06)', 'Kategori (2026/06)', 'Hedef kategori', 'P_kat', 'E_delta', 'Asagi', 'Beklenen Funnel',
        'Surukleyici', 'Risk', 'Not', 'FCF TTM 2026/06', 'FCF statü']].head(80).rename(columns=AD).to_excel(w, sheet_name='KATEGORI_ATLAMA_2026-09', index=False)
    Y.to_excel(w, sheet_name='ADAYLAR_2026-09', index=False)
    MD.to_excel(w, sheet_name='METRIK_DETAY_ILK60', index=False)
    pd.DataFrame(MOD).to_excel(w, sheet_name='MODEL', index=False)
wb = load_workbook('KURULU_YAY.xlsx')
for ws in wb:
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78'); c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 45
    if ws.title == 'METODOLOJI':
        ws.column_dimensions['A'].width = 170
        for r in ws.iter_rows(min_row=2):
            r[0].alignment = Alignment(wrap_text=True, vertical='top')
            if str(r[0].value or '').startswith('#'): r[0].font = Font(bold=True, color='1F4E78', size=12)
        continue
    ws.freeze_panes = 'C2' if ws.title.startswith('ADAY') else 'B2'
    for col in ws.columns:
        h = str(col[0].value)
        ws.column_dimensions[col[0].column_letter].width = 32 if h in ('Ünvan',) or 'sürükleyici' in h or 'riskler' in h else 22 if h in ('Sektör', 'Sınıf', 'Test', 'Yontem') else 12
        fmt = '0%' if (h.startswith('P(') or h.startswith('P ') or 'oran' in h or 'isabet' in h) else '0.00' if h in ('spearman',) or 'σ' in h else '#,##0.0' if any(k in h for k in ('E[Δ]', 'kazanç', 'kayıp', 'Kazanç', 'Kayıp', 'Funnel', 'FCF TTM', 'Δ', 'delta', 'Değer')) else None
        if fmt:
            for c in col[1:]:
                if isinstance(c.value, (int, float)): c.number_format = fmt
    ws.auto_filter.ref = ws.dimensions
for wsn in ('KATEGORI_ATLAMA_2026-09',):
    w2 = wb[wsn]; w2.freeze_panes = 'C2'
    for col in ('L', 'M', 'N'): w2.column_dimensions[col].width = 50
ws = wb['ADAYLAR_2026-09']
for t, clr in (('GÜÇLÜ YAY (dengeli)', 'C6EFCE'), ('YAY — yüksek potansiyel, yüksek risk', 'FFEB9C'), ('ORTA YAY', 'DDEBF7'), ('TERS YAY (düşüş riski)', 'F8CBAD')):
    ws.conditional_formatting.add(f'E2:E{ws.max_row}', CellIsRule(operator='equal', formula=[f'"{t}"'], fill=PatternFill('solid', fgColor=clr)))
for c in ws['E']: c.alignment = Alignment(wrap_text=True)
for col in ('P', 'Q', 'R'):
    ws.column_dimensions[col].width = 55
wb.save('KURULU_YAY.xlsx')
print(S.Sınıf.value_counts())
print(S[['Kod', 'Sınıf', 'E_delta', 'P_sicrama', 'Asagi', 'Funnel MIN (2026/06)', 'Surukleyici', 'Risk']].head(30).round(2).to_string())
