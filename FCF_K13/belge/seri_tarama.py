"""FCF TTM seri taraması (FCF_MASTER.ANA → FCF_SERI_TARAMA.xlsx).
Pozitif seri: son N dönem TTM FCF > 0 (N = 5/4/3; gruplar ayrık: 4'lü listede 5'liye girenler yok, 3'lüde 4'lüye girenler yok).
İyileşen: son 4 dönem TTM FCF her dönem bir öncekinden büyük (kesintisiz artış) ve pencerede en az bir negatif dönem var
(yani 'negatiften daha az negatife / pozitife' giden); son 3 dönemi pozitif olanlar zaten pozitif listelerde.
NULL dönem seriyi keser. KOŞULLU dönem dahil edilir ama işaretlenir (GEM: KOŞULLU rakama dayalı yorum yapılmaz)."""
import pandas as pd
D = ['2025/03', '2025/06', '2025/09', '2025/12', '2026/03', '2026/06']
A = pd.read_excel('FCF_MASTER.xlsx', sheet_name='ANA', header=None)
A.columns = (['Kod', 'Ünvan', 'Kapsam', 'Tip', 'Sektör'] + D + [d + ' st' for d in D] + ['F ' + d for d in D[1:]] + ['Kat ' + d for d in D[1:]])
A = A.iloc[4:].reset_index(drop=True)
A = A[A.Kod.notna()]
for d in D: A[d] = pd.to_numeric(A[d], errors='coerce')
P = pd.read_excel('PFCF_GETIRI.xlsx', sheet_name='P_FCF', header=None)
h = next(i for i in range(10) if 'Kod' in [str(x) for x in P.iloc[i]])
P.columns = P.iloc[h]; P = P.iloc[h + 1:]; P = P[P.Kod.notna()]
pc = [c for c in P.columns if isinstance(c, str) and (c.startswith('Güncel') or c.startswith('Piyasa') or c.startswith('P/FCF') or 'getiri' in c.lower() or '%' in c)]
A = A.merge(P[['Kod'] + pc], on='Kod', how='left')

def pozitif(r, n): return all(pd.notna(r[d]) and r[d] > 0 for d in D[-n:])
def artan(r, n):
    v = [r[d] for d in D[-n:]]
    return all(pd.notna(x) for x in v) and all(b > a for a, b in zip(v, v[1:]))
def kos(r, n): return ', '.join(d for d in D[-n:] if r[d + ' st'] == 'KOŞULLU')

A['P5'] = A.apply(lambda r: pozitif(r, 5), axis=1)
A['P4'] = A.apply(lambda r: pozitif(r, 4), axis=1) & ~A.P5
A['P3'] = A.apply(lambda r: pozitif(r, 3), axis=1) & ~A.P5 & ~A.P4
A['6/6 pozitif'] = A.apply(lambda r: 'EVET' if pozitif(r, 6) else '', axis=1)
ust = A.apply(lambda r: pozitif(r, 3), axis=1)
A['IY'] = A.apply(lambda r: artan(r, 4) and min(r[d] for d in D[-4:]) < 0, axis=1) & ~ust
A['Artış 5 dönem kesintisiz'] = A.apply(lambda r: 'EVET' if artan(r, 5) else '', axis=1)
A['Son dönem işareti'] = A['2026/06'].map(lambda x: 'pozitif' if x > 0 else 'negatif' if x < 0 else '')
A['4 dönemde iyileşme (mn TL)'] = A['2026/06'] - A['2025/09']
A['Son kategori'] = A['Kat 2026/06']

def liste(m, n):
    x = A[m].copy()
    x['KOŞULLU dönem'] = x.apply(lambda r: kos(r, n), axis=1)
    x['Son dönem statü'] = x['2026/06 st']
    return x

KOL = ['Kod', 'Ünvan', 'Tip', 'Sektör'] + D + ['Son dönem statü', 'KOŞULLU dönem'] + pc + ['Son kategori']
S = {'POZITIF_5': (liste(A.P5, 5), KOL[:4] + ['6/6 pozitif'] + KOL[4:]),
     'POZITIF_4': (liste(A.P4, 4), KOL), 'POZITIF_3': (liste(A.P3, 3), KOL),
     'IYILESEN': (liste(A.IY, 4), KOL[:10] + ['Son dönem işareti', '4 dönemde iyileşme (mn TL)', 'Artış 5 dönem kesintisiz'] + KOL[10:])}
OK = [('Kaynak', 'FCF_MASTER.xlsx (ANA): FCF TTM, mn TL, Haziran 2026 satın alma gücü. Fiyat/P/FCF/getiri: PFCF_GETIRI.xlsx (06.10.2026 ~16:40).'),
      ('POZITIF_5', 'Son 5 dönem (2025/06–2026/06) TTM FCF > 0. "6/6 pozitif" = 2025/03 da pozitif.'),
      ('POZITIF_4', 'Son 4 dönem (2025/09–2026/06) pozitif, 2025/06 pozitif değil (veya NULL). 5\'li listedekiler burada yok.'),
      ('POZITIF_3', 'Son 3 dönem (2025/12–2026/06) pozitif, 2025/09 pozitif değil (veya NULL). 4\'lü/5\'li listedekiler burada yok.'),
      ('IYILESEN', 'Son 4 dönem (2025/09 → 2026/06) TTM FCF her dönem bir öncekinden yüksek (kesintisiz artış) ve pencerede en az bir negatif '
                   'dönem var; son 3 dönemi zaten pozitif olanlar pozitif listelerde. "Son dönem işareti" pozitifse şirket negatiften pozitife geçmiş.'),
      ('Statü', 'NULL dönem seriyi keser. KOŞULLU dönem dahil edildi ama "KOŞULLU dönem" sütununda işaretli (GEM: KOŞULLU rakama dayalı yorum yapılmaz).'),
      ('Not', 'TTM serisi her çeyrek bir çeyrek kayar; ardışık iki TTM dönemi üç çeyreği paylaşır. Kapsam dışı (banka, sigorta, GYO vb.) şirketlerin FCF\'i yok.')]
with pd.ExcelWriter('FCF_SERI_TARAMA.xlsx', engine='openpyxl') as w:
    pd.DataFrame(OK, columns=['Konu', 'Açıklama']).to_excel(w, sheet_name='OKUBENI', index=False)
    ozet = [(k, len(v[0])) for k, v in S.items()]
    pd.DataFrame(ozet, columns=['Liste', 'Şirket']).to_excel(w, sheet_name='OZET', index=False)
    for k, (x, kol) in S.items():
        x = x.sort_values('2026/06' if k != 'IYILESEN' else '4 dönemde iyileşme (mn TL)', ascending=False)
        x[kol].to_excel(w, sheet_name=k, index=False)
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
wb = load_workbook('FCF_SERI_TARAMA.xlsx')
for ws in wb:
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78'); c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 40
    if ws.title == 'OKUBENI':
        ws.column_dimensions['A'].width = 14; ws.column_dimensions['B'].width = 140
        for r in ws.iter_rows(min_row=2):
            r[1].alignment = Alignment(wrap_text=True, vertical='top')
        continue
    ws.freeze_panes = 'C2'
    for col in ws.columns:
        h = str(col[0].value)
        ws.column_dimensions[col[0].column_letter].width = 34 if h == 'Ünvan' else 22 if h in ('Sektör', 'Son kategori') else 12
        if h in D or 'iyileşme' in h or h.startswith('Piyasa'):
            for c in col[1:]: c.number_format = '#,##0.0;[Red]-#,##0.0'
        if h.startswith('P/FCF'):
            for c in col[1:]: c.number_format = '0.00'
        if '%' in h or 'getiri' in h.lower():
            for c in col[1:]: c.number_format = '0.0%' if isinstance(c.value, float) and abs(c.value) < 20 else c.number_format
    ws.auto_filter.ref = ws.dimensions
wb.save('FCF_SERI_TARAMA.xlsx')
print({k: len(v[0]) for k, v in S.items()}, pc)
