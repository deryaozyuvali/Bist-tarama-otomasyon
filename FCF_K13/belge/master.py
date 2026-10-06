"""FCF + Funnel MASTER dosyası (GEM v6.2).
Tek kaynak: FCF TTM (bileşenleri formüllü) · belgeden kurulan TTM parçaları · rapor kalemleri (belgeler) · elle kararlar ·
notlar · Funnel V8.2 (metrik + puan, M1 güncel CFO, Bölüm F) · kod listeleri. Yeni dönem: satır eklenir (OKUBENI).
Kullanım (FCF_K13 içinden): python3 belge/master.py <GEM_20xxQx_Funnel_V8_2_FCF_V2_1.xlsx ...>   Çıktı: FCF_MASTER.xlsx"""
import sys, re, json, ast
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

DON = ['2025/03', '2025/06', '2025/09', '2025/12', '2026/03', '2026/06']
FDON = ['2025/06', '2025/09', '2025/12', '2026/03', '2026/06']
Q = {'Q1': '03', 'Q2': '06', 'Q3': '09', 'Q4': '12'}
PUAN = [('M1', 15), ('M2', 10), ('M3', 5), ('M5', 12), ('M6', 11), ('M7', 10), ('M8', 10), ('M9', 8), ('M10', 7), ('M11', 7), ('M12', 5)]
F_ = Font(name='Arial', size=10); FB = Font(name='Arial', size=10, bold=True); FH = Font(name='Arial', size=10, bold=True, color='FFFFFF')
FT = Font(name='Arial', size=14, bold=True); FBL = Font(name='Arial', size=10, color='0000FF'); FG = Font(name='Arial', size=10, color='008000')
FI = Font(name='Arial', size=9, italic=True)
HDR = PatternFill('solid', fgColor='1F3864'); GRP = PatternFill('solid', fgColor='D9E1F2')
NUM = '#,##0.0;[Red]-#,##0.0;-'; TL = '#,##0;[Red]-#,##0;-'; RAT = '0.00;[Red]-0.00'


def bos(x):
    return x is None or x == '' or (isinstance(x, float) and pd.isna(x))


def temiz(x):
    if isinstance(x, str): return ILLEGAL_CHARACTERS_RE.sub('', x)
    return None if bos(x) else x


def sayi(x):
    try:
        return None if bos(x) else float(x)
    except ValueError:
        return None


def ast_dict(dosya, ad):
    for n in ast.parse(open(dosya).read()).body:
        if isinstance(n, ast.Assign) and getattr(n.targets[0], 'id', '') == ad:
            return eval(compile(ast.Expression(n.value), dosya, 'eval'), {})
    return {}


wb = Workbook()


def sayfa(ad, kolonlar, satirlar, genis, bicim=None, font=None, dondur='B2', idx=None):
    ws = wb.create_sheet(ad) if idx is None else wb.create_sheet(ad, idx)
    for j, c in enumerate(kolonlar, 1):
        x = ws.cell(1, j, c); x.font = FH; x.fill = HDR; x.alignment = Alignment(wrap_text=True, vertical='center')
    for i, r in enumerate(satirlar, 2):
        for j, v in enumerate(r, 1):
            c = ws.cell(i, j, temiz(v)); c.font = (font or {}).get(j, F_)
            if bicim and j in bicim: c.number_format = bicim[j]
    for j, w in enumerate(genis, 1): ws.column_dimensions[L(j)].width = w
    ws.freeze_panes = dondur; ws.auto_filter.ref = f'A1:{L(len(kolonlar))}{len(satirlar) + 1}'
    ws.row_dimensions[1].height = 45
    return ws


# =============== FCF verisi ===============
N = pd.read_excel('FCF_TTM_nihai_belgeli.xlsx', sheet_name='FCF_TTM').sort_values(['Kod', 'Dönem'], ignore_index=True)
N['Statü'] = N['Statü'].fillna('NULL')  # pandas 'NULL' metnini NaN okur
nul = pd.read_csv('belge/null_ttm.csv', keep_default_na=False)
kos = pd.read_csv('belge/kosullu_sonuc.csv', keep_default_na=False)
hol = pd.read_csv('belge/holding_ttm.csv', keep_default_na=False)
k6t = pd.read_csv('belge/k6_ttm.csv', keep_default_na=False)
K6I = k6t.set_index(['Kod', 'Dönem'])
NUL, KOS = nul.set_index(['Kod', 'Dönem']), kos.set_index(['Kod', 'Dönem'])
ACIK = ast_dict('belge/teslim_kosullu.py', 'ACIKLAMA'); NOTD = ast_dict('belge/teslim_kosullu.py', 'NOT')
ESKI_KOD = ast_dict('belge/nihai.py', 'ESKI_KOD')

# CAPEX ayrımı (MDV+MODV / YAGM)
B = pd.read_excel('girdi/FCF_TTM_2025-03_2026-06_v7_nihai.xlsx', sheet_name='Bileşenler')
B = B[B['Bileşen'].isin(['MDV+MODV alımı', 'YAGM alımı'])].pivot_table(index=['Kod', 'Dönem'], columns='Bileşen',
                                                                        values='Formülde (mn TL)', aggfunc='first')
def kul(x):
    x = sayi(x); return None if x is None else (-x if x < 0 else 0.0)  # K6: pozitif TTM alım/kira satırı kullanılmaz
BEL = {}
for d in (nul, kos, k6t):
    for _, r in d.iterrows(): BEL[(r['Kod'], r['Dönem'])] = (kul(r['MDV+MODV TTM']), kul(r['YAGM TTM']))
mdv, yagm = [], []
for _, r in N.iterrows():
    k, cap = (r['Kod'], r['Dönem']), r['CAPEX_STD']
    if pd.isna(cap): mdv.append(None); yagm.append(None); continue
    cand = []
    if 'BELGE' in str(r['Kaynak']) or 'teyitli' in str(r['Kaynak']):
        if k in BEL and None not in BEL[k]: cand.append(BEL[k])
    if k in B.index:
        a, b = B.loc[k].get('MDV+MODV alımı'), B.loc[k].get('YAGM alımı')
        cand.append((0.0 if pd.isna(a) else abs(a), 0.0 if pd.isna(b) else abs(b)))
    ok = next((c for c in cand if abs(c[0] + c[1] - cap) < 0.01), None) or ((cap - (cand[0][1] if cand else 0.0)), cand[0][1] if cand else 0.0)
    mdv.append(ok[0]); yagm.append(ok[1])
N['MDV+MODV'] = mdv; N['YAGM'] = yagm

# GEM v6.2 kodları
def null_kodu(kod, don, st):
    if st != 'NULL': return None, None
    if (kod, don) in K6I.index and (kod, don) not in NUL.index:
        return ('NULL_BELGE_YOK' if K6I.loc[(kod, don)]['Satır statü'] == 'BELGE_YOK' else 'NULL_OKUNAMADI'), ''
    if (kod, don) not in NUL.index: return 'NULL_BELGE_YOK', ''
    b = NUL.loc[(kod, don)]
    if b['Satır statü'] not in ('BELGE_YOK',): return 'NULL_OKUNAMADI', ''
    metin = ' ‖ '.join(b[k + ' kanıt'] for k in ('CFO', 'MDV+MODV', 'YAGM', 'KIRA'))
    if 'MALI_YIL' in metin: return 'NULL_BELGE_YOK', 'Mali yıl takvim yılı değil; takvim TTM\'i şirket raporlarından kurulamıyor'
    eksik = set(re.findall(r'(\d{4}/\d\d) raporu KAP', metin))
    if eksik and all(e < don for e in eksik): return 'NULL_GECMIS_YOK', 'Önceki dönem raporu yok (' + ', '.join(sorted(eksik)) + ')'
    return 'NULL_BELGE_YOK', ''

KIYAS = {'CFO': 'CFO kıyas', 'CAPEX': 'CAPEX kıyas', 'KIRA': 'KIRA kıyas'}
def fark_kodu(kod, don, kaynak='', statu='', aciklama=''):
    if (kod, don) in K6I.index and (kod, don) not in KOS.index:
        if 'teyit' in str(kaynak): return 'TMS29_KATSAYI_ACIKLIYOR (K6/K7 belgeyle teyit)'
        if 'BELGE' in str(kaynak) or 'KARMA' in str(kaynak): return 'EVO_ESLEME_HATASI (K6/K7)'
        if 'maddi işaret' in str(aciklama): return 'ISARET_ANOMALISI_COZULEMEDI (K6)'
        return None
    if (kod, don) not in KOS.index: return None
    b = KOS.loc[(kod, don)]
    if b['Sonuç'] == 'BELGE_YOK': return 'BELGE_YOK'
    if b['Sonuç'] == 'ELLE': return 'ACIKLANAMAYAN_FARK'
    if f'{kod}|{don}' in NOTD and kod == 'SAHOL': return 'YENIDEN_DUZENLEME (K11): CFO, CAPEX'
    parca = []
    for k, c in KIYAS.items():
        v = b[c]
        kan = b['CFO kanıt'] if k == 'CFO' else b['MDV+MODV kanıt'] if k == 'CAPEX' else b['KIRA kanıt']
        if v == 'FARKLI': parca.append(f'{k}: EVO_ESLEME_HATASI')
        elif v == 'EVO_NOMINAL': parca.append(f'{k}: EVO_YENIDEN_BAZLAMAMIS')
        elif v == 'YAKIN': parca.append(f'{k}: YAKIN')
    if parca: return '; '.join(parca)
    kat = re.search(r'TMS29 ×(?!1\.0000)', b['CFO kanıt'] + b['MDV+MODV kanıt'] + b['KIRA kanıt'])
    return 'TMS29_KATSAYI_ACIKLIYOR' if kat else 'BIREBIR'

N['NULL kodu'], N['NULL notu'] = zip(*[null_kodu(k, d, s) for k, d, s in zip(N.Kod, N['Dönem'], N['Statü'])])
N['Evo–belge kodu (14g)'] = [fark_kodu(k, d, s, st, a) for k, d, s, st, a in zip(N.Kod, N['Dönem'], N['Kaynak'], N['Statü'], N['Açıklama'])]
def aciklama(r):
    a = r['Açıklama'] if isinstance(r['Açıklama'], str) else ''
    n = NOTD.get(f"{r['Kod']}|{r['Dönem']}")
    if isinstance(r['NULL notu'], str) and r['NULL notu']: a = r['NULL notu'] + '. ' + a
    return (a + (' · NOT: ' + n if n else '')).strip()
N['Açıklama2'] = N.apply(aciklama, axis=1)

# =============== FCF sayfası ===============
FC = ['Anahtar', 'Kod', 'Tip', 'Sektör', 'Dönem', 'CFO_TTM', '|MDV+MODV alımı|', '|YAGM alımı|', 'CAPEX_STD', '|Kira anapara|',
      'Alınan temettü (CFO dışı)', 'FCF_STD (①)', 'FCF_HLD (②)', 'FCF türü', 'FCF · TTM (mn TL)', 'Statü', 'Kaynak', 'NULL kodu',
      'Evo–belge kodu (14g)', 'Açıklama', 'V7 FCF', 'V7 statü', "V7'ye göre değişti", 'Belge FCF (çözülmemiş çatışma)']
wf = wb.active; wf.title = 'FCF'
for j, c in enumerate(FC, 1):
    x = wf.cell(1, j, c); x.font = FH; x.fill = HDR; x.alignment = Alignment(wrap_text=True, vertical='center')
FORMUL = {9: 'CAPEX_STD = |MDV+MODV alımı| + |YAGM alımı| (GEM FCF v2.2)', 12: 'FCF_STD = CFO_TTM − CAPEX_STD − |Kira anapara|',
          13: 'FCF_HLD = FCF_STD + Alınan temettü (yatırım bölümü, CFO dışı) — yalnız Holding tipinde',
          15: 'Ana FCF: FCF türü FCF_HLD ise FCF_HLD, değilse FCF_STD', 6: 'Mavi = Evo girdisi; yeşil = BELGE_TTM sayfasından formülle gelen belge TTM\'i (A×kA + B×kB − C×kC).'}
for j, t in FORMUL.items(): wf.cell(1, j).comment = Comment(t, 'GEM v6.2')
for i, (_, r) in enumerate(N.iterrows(), 2):
    v = lambda c: temiz(r[c])
    vals = {1: f"{r['Kod']}|{r['Dönem']}", 2: r['Kod'], 3: v('Tip'), 4: v('Sektör'), 5: r['Dönem'], 6: v('CFO_TTM'), 7: v('MDV+MODV'),
            8: v('YAGM'), 9: f'=IF(OR(F{i}="",G{i}="",H{i}=""),"",G{i}+H{i})', 10: v('|Kira anapara|'), 11: v('Alınan temettü (CFO dışı)'),
            12: f'=IF(OR(F{i}="",I{i}="",J{i}=""),"",F{i}-I{i}-J{i})', 13: f'=IF(OR(L{i}="",N{i}<>"FCF_HLD"),"",L{i}+N(K{i}))',
            14: v('FCF türü'), 15: f'=IF(L{i}="","",IF(N{i}="FCF_HLD",M{i},L{i}))', 16: v('Statü'), 17: v('Kaynak'),
            18: v('NULL kodu'), 19: v('Evo–belge kodu (14g)'), 20: v('Açıklama2'), 21: v('V7 FCF'), 22: v('V7 statü'),
            23: v("V7'ye göre değişti"), 24: v('Belge FCF (çözülmemiş)') if 'Belge FCF (çözülmemiş)' in r else None}
    for j, x in vals.items():
        c = wf.cell(i, j, x); c.font = FBL if j in (6, 7, 8, 10, 11) else F_
        if j in (6, 7, 8, 9, 10, 11, 12, 13, 15, 21, 24): c.number_format = NUM
nF = len(N) + 1
wf.freeze_panes = 'F2'; wf.auto_filter.ref = f'A1:X{nF}'; wf.column_dimensions['A'].hidden = True
for j, w in enumerate([14, 8, 10, 22, 9] + [12] * 8 + [10, 13, 11, 22, 17, 30, 70, 11, 11, 14, 14], 1): wf.column_dimensions[L(j)].width = w
wf.row_dimensions[1].height = 45

# =============== BELGE_TTM (TTM parçaları) ===============
KAL = ['CFO', 'MDV+MODV', 'YAGM', 'KIRA', 'TEMETTU']
bt = []
kullanilan = {(r['Kod'], r['Dönem']) for _, r in N.iterrows() if 'BELGE' in str(r['Kaynak']) or 'teyitli' in str(r['Kaynak'])}
K6SET = set(K6I.index)
for grup, d in (('NULL', nul), ('KOŞULLU', kos), ('HOLDING (temettü)', hol), ('K6/K7', k6t)):
    for _, r in d.iterrows():
        for k in KAL:
            if k + ' bileşen' not in r or (grup.startswith('HOLDING') and k != 'TEMETTU'): continue
            bil = r[k + ' bileşen']
            if bil == '' and r[k + ' statü'] == '': continue
            p = [sayi(x.strip().replace('—', '')) for x in bil.split('|')] if bil else [None] * 3
            if len(p) == 1 and p[0] is not None: p = [p[0], 0.0, 0.0]  # yıl sonu: TTM = FY
            p = (p + [None] * 3)[:3]
            kan = [s.strip() for s in r[k + ' kanıt'].split('‖')]; kan = (kan + [''] * 3)[:3]
            if r[k + ' statü'] == 'IZAHNAME':  # izahname: parçalar raporda yazan (nominal), tek katsayı (izahname tarihi)
                m_ = re.search(r'×([\d.]+)', kan[0]); kat = [sayi(m_.group(1)) if m_ else 1.0] * 3
            else:  # KAP/site/Evo PDF: bileşen zaten çevrilmiş → raporda yazan = çevrilmiş / katsayı
                kat = [sayi(m_.group(1)) if (m_ := re.search(r'TMS29 ×([\d.]+)', s)) else 1.0 for s in kan]
                p = [None if x is None else round(x / c, 3) for x, c in zip(p, kat)]
            bt.append([f"{r['Kod']}|{r['Dönem']}|{k}", r['Kod'], r['Dönem'], k, grup, *p, None, None, *kat, r[k + ' statü'],
                       'EVET' if (r['Kod'], r['Dönem']) in kullanilan and grup != 'HOLDING (temettü)' else '', *kan])
BTC = ['Anahtar', 'Kod', 'Dönem', 'Kalem', 'Grup', 'A: cari dönem YTD (raporda, mn TL)', 'B: önceki yıl sonu FY (raporda, mn TL)',
       'C: önceki yıl aynı dönem YTD (raporda, mn TL)', 'TTM = A×kA + B×kB − C×kC', 'Formülde kullanılan', 'Katsayı A', 'Katsayı B', 'Katsayı C', 'Kalem statüsü', "FCF'te kullanıldı",
       'Kanıt A', 'Kanıt B', 'Kanıt C']
wt = sayfa('BELGE_TTM', BTC, bt, [16, 8, 9, 10, 14, 13, 13, 13, 13, 13, 9, 9, 9, 13, 9, 60, 60, 60],
           bicim={6: NUM, 7: NUM, 8: NUM, 11: '0.0000', 12: '0.0000', 13: '0.0000'}, font={6: FBL, 7: FBL, 8: FBL}, dondur='F2')
for i in range(2, len(bt) + 2):
    wt.cell(i, 9, f'=IF(OR(F{i}="",G{i}="",H{i}=""),"",F{i}*K{i}+G{i}*L{i}-H{i}*M{i})').number_format = NUM
    wt.cell(i, 10, f'=IF(I{i}="","",IF(D{i}="CFO",I{i},IF(D{i}="TEMETTU",MAX(I{i},0),IF(I{i}>0,0,-I{i}))))').number_format = NUM
    wt.cell(i, 9).font = F_; wt.cell(i, 10).font = F_
wt.cell(1, 6).comment = Comment('Parçalar raporda yazan tutar (mn TL, çevrilmemiş). k = TMS 29 katsayısı → Haziran 2026 (TMS29 sayfası). Yıl sonunda TTM = A. Ham TL tutar RAPOR_KALEM sayfasında.', 'GEM v6.2')
wt.cell(1, 10).comment = Comment('CFO: TTM. Alım/kira: |TTM|; TTM pozitifse 0 (Kontrol 6). Temettü: pozitif kısım.', 'GEM v6.2')
wt.column_dimensions['A'].hidden = True

# FCF girdilerini BELGE_TTM'e bağla (belgeden gelen satırlarda CFO / MDV+MODV / YAGM / KIRA = BELGE_TTM "Formülde kullanılan")
BIDX = {}
for j, b in enumerate(bt, 2):
    if b[14] != 'EVET' or b[3] == 'TEMETTU': continue
    a_, b_, c_, ka, kb, kc = b[5], b[6], b[7], b[10], b[11], b[12]
    if None in (a_, b_, c_, ka, kb, kc): continue
    ttm = a_ * ka + b_ * kb - c_ * kc
    BIDX[b[0]] = (j, ttm if b[3] == 'CFO' else (0.0 if ttm > 0 else -ttm))
bag = 0
for i, (_, r) in enumerate(N.iterrows(), 2):
    if not ('BELGE' in str(r['Kaynak']) or 'teyitli' in str(r['Kaynak'])): continue
    for kal, colj, val in (('CFO', 6, r['CFO_TTM']), ('MDV+MODV', 7, r['MDV+MODV']), ('YAGM', 8, r['YAGM']), ('KIRA', 10, r['|Kira anapara|'])):
        k = f"{r['Kod']}|{r['Dönem']}|{kal}"
        if k in BIDX and not bos(val) and abs(BIDX[k][1] - val) < 0.01:
            c = wf.cell(i, colj, f'=BELGE_TTM!J{BIDX[k][0]}'); c.font = FG; bag += 1
print('FCF girdisi BELGE_TTM bağlantısı:', bag)

# =============== RAPOR_KALEM (belgeler) ===============
rk = []
for f in ('belge/rapor_kalemleri.csv', 'belge/rapor_kalemleri_site.csv', 'belge/rapor_kalemleri_evo.csv'):
    d = pd.read_csv(f, keep_default_na=False)
    for _, r in d.iterrows():
        idx = str(r['idx']).split('.')[0]
        bag = f'https://www.kap.org.tr/tr/Bildirim/{idx}' if idx.isdigit() else ''
        kay = r['kaynak']
        if 'http' in kay: bag = kay[kay.index('http'):]; kay = kay[:kay.index('http')].strip()
        rk.append([r['kod'], f"{r['yil']}/{int(r['ay']):02d}", r['kalem'], r['sutun'], sayi(r['deger_tl']), kay, r['durum'], idx, bag,
                   r['pdf_birim'], r['pdf_sayfa'], r['pdf_satir'], r['tms29'], r['not_']])
iz = pd.read_csv('belge/izahname_kalemleri.csv', keep_default_na=False)
for _, r in iz.iterrows():
    rk.append([r['kod'], f"{r['yil']}/{int(r['ay']):02d}", r['kalem'], 'izahname', sayi(r['deger_tl']), 'İzahname: ' + r['ek'], r['durum'],
               str(r['idx']), f"https://www.kap.org.tr/tr/Bildirim/{r['idx']}", 'TL', r['sayfa'], r['satir'], '', ''])
RKC = ['Kod', 'Rapor dönemi', 'Kalem', 'Sütun (cari / önceki)', 'Değer (TL, raporda yazan)', 'Kaynak', 'Okuma durumu', 'KAP bildirim no',
       'Bağlantı', 'Birim', 'Sayfa', 'Satır (belgedeki metin)', 'TMS 29 başlığı', 'Not']
sayfa('RAPOR_KALEM', RKC, rk, [8, 10, 10, 10, 18, 20, 13, 11, 40, 7, 6, 70, 8, 50], bicim={5: TL}, dondur='C2')

# =============== ELLE_KARAR ===============
ek = json.load(open('belge/elle_kararlar.json'))
er = []
for k, v in ek.items():
    if k.startswith('_'): continue
    p = k.split('|')
    v = v if isinstance(v, dict) else {'deger': v}
    er.append([p[0], f'{p[1]}/{int(p[2]):02d}' if len(p) > 2 else '', p[3] if len(p) > 3 else '', p[4] if len(p) > 4 else '',
               sayi(v.get('deger')) if not isinstance(v.get('deger'), str) else None, v.get('durum', ''), v.get('not', '')])
sayfa('ELLE_KARAR', ['Kod', 'Rapor dönemi', 'Kalem', 'Sütun', 'Belgedeki tutar (TL)', 'Durum', 'Gerekçe'], er,
      [8, 10, 10, 8, 18, 12, 120], bicim={5: TL}, dondur='C2')

# =============== NOTLAR ===============
nt = [[k, '', 'ELLE (belge çözemedi)', v] for k, v in ACIK.items()]
nt += [[k.split('|')[0], k.split('|')[1], 'Belge kararı', v] for k, v in NOTD.items()]
nt += [[k, '', 'Kod değişikliği', f'Eski kod → {v}' if v else 'Kod hiçbir kaynakta yok'] for k, v in ESKI_KOD.items()]
GENEL = [
    ('Restatement (yeniden düzenleme)', 'TTM\'in üç parçası aynı esasta olmalı; esas, cari raporun karşılaştırmalı sütunundan anlaşılır. '
     'SAHOL: 2024 sınıflaması ilk kez 2025/12 raporunda değişti → 2025 ara dönemlerinde orijinal 2024 esası. KATMR: esas belirlenemedi.'),
    ('XBRL etiket hataları', 'Satış satırı alım diye etiketlenmiş (BRLSM, DENGE, MEGMT) · avans netleşmiş (CIMSA) · kullanım hakkı alımı '
     'MDV\'ye katılmış (ARZUM, HUNER) · yanlış CFO (BOBET 2025/06, HKTM, SNICA, TUREX 2024). Karar: PDF satırı + özdeşlik.'),
    ('Evo yeniden bazlama istisnası', 'MEGMT ve bazı kira satırlarında Evo rakamı belgenin çevrilmemiş hâli; BLCYT/GMTAS/ECOGR 2025/12 '
     'Haziran 2026 bazına taşınmamış.'),
    ('Kaynak sırası', 'KAP bildirimi (XBRL + PDF) → şirket sitesi (Yatırımcı İlişkileri) → Evo belge havuzu. Evo sayısal verisi yalnız kıyas.'),
    ('Özdeşlik', 'İşletme + Yatırım + Finansman (+ kur farkı etkisi) = nakit değişimi; XBRL ile PDF farklıysa özdeşliği sağlayan alınır.'),
    ('Kontrol 9', 'Yatırım bölümü "Diğer" artığı büyük satırlar bileşenleri belgeyle doğrulansa da KOŞULLU (dipnot sınıflaması gerekir).'),
    ('Spor kulüpleri', 'BJKAS, FENER, GSRAY, TSPOR, KAYSE, MERKO: mali yıl Haziran/Mayıs; takvim TTM\'i yok (NULL).'),
    ('Eski taramayla karşılaştırma', 'ESKI_TARAMA_karsilastirma.xlsx: 1.654/1.886 satır tutuyor; 163 satır belge düzeltmesi, 69 satır '
     'kural farkı (ör. ENJSA imtiyaz yatırımı capex sayıldı).'),
]
nt += [['GENEL', '', k, v] for k, v in GENEL]
sayfa('NOTLAR', ['Kod', 'Dönem', 'Tür', 'Not'], nt, [10, 9, 26, 140], dondur='A2')

# =============== TMS29 ===============
TM = [('2024/12', 1.5414), ('2025/03', 1.4004), ('2025/06', 1.3211), ('2025/09', 1.2289), ('2025/12', 1.1776), ('2026/03', 1.0701), ('2026/06', 1.0)]
ws = sayfa('TMS29', ['Rapor tarihi', 'Katsayı → Haziran 2026', 'Açıklama'], [[d, k, ''] for d, k in TM], [12, 18, 100], bicim={2: '0.0000'}, font={2: FBL})
ws['C2'] = ('Belge rakamı × katsayı = Haziran 2026 satın alma gücü. Kaynak: Evo/PDF oranlarından ölçüldü (BIMAS 21/21 kalem birebir; '
            'şüpheli 405 satırda CFO\'da 282 birebir). GEM v6.2 Kontrol 14 (g): katsayı farkı AÇIKLAR, Evo yöntemini doğrulamaz; '
            'mümkünse şirketin kendi raporundaki TMS 29 endeksi kullanılır. Şirketin son raporu 2026/06 değilse bölen o tarihin katsayısıdır.')
ws['C2'].alignment = Alignment(wrap_text=True, vertical='top'); ws['C2'].font = F_

# =============== FUNNEL ===============
FD, FG_ = [], []
for f in sys.argv[1:]:
    y, q = re.search(r'GEM_(\d{4})(Q\d)', f).groups(); et = f'{y}/{Q[q]}'
    son = pd.Timestamp(f'{y}-{Q[q]}-01') + pd.offsets.MonthEnd(0)
    d = pd.read_excel(f, sheet_name='Funnel_Detay', header=2); d = d[d.Kod.notna()]
    g = pd.read_excel(f, sheet_name='Evo_Girdileri', header=2).rename(columns={'hisse_senedi_kodu': 'Kod'}); g = g[g.Kod.notna()]
    def takvim(yay):
        yay = pd.Timestamp(yay)
        if pd.isna(yay) or (yay - son).days <= 110: return et
        x = yay - pd.Timedelta(days=20)
        qe = pd.Timestamp(x.year, ((x.month - 1) // 3) * 3 + 1, 1) - pd.Timedelta(days=1)
        return f'{qe.year}/{qe.month:02d}'
    g['Dönem'] = g['yayinlanma_tarihi_utc'].map(takvim); g['Dosya etiketi'] = et
    d = d.merge(g[['Kod', 'Dönem', 'Dosya etiketi', 'cfo_ttm', 'net_total_cur']], on='Kod', how='left')
    d['Dönem'] = d['Dönem'].fillna(et); d['Dosya etiketi'] = d['Dosya etiketi'].fillna(et)
    FD.append(d); FG_.append(g)
FD = pd.concat(FD, ignore_index=True).drop_duplicates(['Kod', 'Dönem'], keep='last').sort_values(['Kod', 'Dönem'])
FG_ = pd.concat(FG_, ignore_index=True).drop_duplicates(['Kod', 'Dönem'], keep='last')
NC = N.set_index(['Kod', 'Dönem'])
MET = ['M1 CFO/NK', 'M2 Net Borç/FAVÖK', 'M3 Cari Oran', 'M5 ROE', 'M6 Brüt Marj', 'M7 Satış Büyümesi', 'M8 Özkaynak Büyümesi',
       'M9 OPEX/Brüt Kâr', 'M10 R', 'M11 FAVÖK/NK', 'M12 Marj Trendi']
DC = (['Anahtar', 'Kod', 'Ünvan', 'Dönem', 'Dosya etiketi', 'Kapsam', 'Kritik Eksik #', 'Puan Durumu', 'CFO TTM (mn TL)', 'CFO kaynağı',
       'Net kâr TTM (mn TL)', 'Eski CFO (Evo, mn TL)'] + sum([[m, m.split()[0] + ' Puan'] for m in MET], []) +
      ['NULL Metrik #', 'Funnel MIN', 'Funnel MAX', 'Tier MIN', 'Tier MAX', 'Kategori', 'Önceki dönem', '2 önceki dönem',
       'FCF son 3 çeyrek (TTM)', 'Bölüm F', 'Eski M1 Puan', 'Eski Funnel MIN', 'Eski Funnel MAX', 'Eski Kategori', 'Değişiklik', 'Not'])
wd = wb.create_sheet('FUNNEL')
for j, c in enumerate(DC, 1):
    x = wd.cell(1, j, c); x.font = FH; x.fill = HDR; x.alignment = Alignment(wrap_text=True, vertical='center')
ci = {c: j for j, c in enumerate(DC, 1)}; col = lambda c: L(ci[c])
PC = [col(m.split()[0] + ' Puan') for m in MET]
FK = lambda c, key: f'INDEX(FCF!${c}$2:${c}${nF},MATCH({key},FCF!$A$2:$A${nF},0))'
for i, (_, r) in enumerate(FD.iterrows(), 2):
    v = lambda c: None if c not in r or bos(r[c]) else r[c]
    kod, don = r['Kod'], r['Dönem']
    ecfo = None if pd.isna(r['cfo_ttm']) else r['cfo_ttm'] / 1e6
    ncfo, kay = ecfo, 'Evo (eski tarama)'
    if (kod, don) in NC.index and not pd.isna(NC.loc[(kod, don), 'CFO_TTM']):
        x = NC.loc[(kod, don)]
        if ecfo is None or abs(x['CFO_TTM'] - ecfo) > max(0.005 * abs(x['CFO_TTM']), 0.5):
            ncfo, kay = x['CFO_TTM'], 'FCF sayfası (' + str(x['Kaynak']) + ')'
    if (kod, don) in KOS.index and KOS.loc[(kod, don)]['Sonuç'] == 'ELLE' and KOS.loc[(kod, don)]['CFO kıyas'] == 'FARKLI':
        ncfo, kay = None, 'NULL_CATISMA: Evo ≠ belge, açıklanamadı (GEM v6.2 A3) → M1 NULL'
    vals = {'Anahtar': f'{kod}|{don}', 'Kod': kod, 'Ünvan': v('Ünvan'), 'Dönem': don, 'Dosya etiketi': v('Dosya etiketi'),
            'Kapsam': v('Kapsam'), 'Kritik Eksik #': v('Kritik Eksik #'), 'Puan Durumu': v('Puan Durumu'), 'CFO TTM (mn TL)': ncfo,
            'CFO kaynağı': kay, 'Net kâr TTM (mn TL)': None if pd.isna(r['net_total_cur']) else r['net_total_cur'] / 1e6,
            'Eski CFO (Evo, mn TL)': ecfo, 'Eski M1 Puan': v('M1 Puan'), 'Eski Funnel MIN': v('Funnel MIN'),
            'Eski Funnel MAX': v('Funnel MAX'), 'Eski Kategori': v('Son Kategori'), 'Not': v('Not')}
    for m in MET[1:]: vals[m] = v(m); vals[m.split()[0] + ' Puan'] = v(m.split()[0] + ' Puan')
    for c, x in vals.items(): wd.cell(i, ci[c], temiz(x))
    C, NK, D_ = f"{col('CFO TTM (mn TL)')}{i}", f"{col('Net kâr TTM (mn TL)')}{i}", f"{col('Dönem')}{i}"
    P1, P2 = f"{col('Önceki dönem')}{i}", f"{col('2 önceki dönem')}{i}"
    onc = lambda d: f'IF(RIGHT({d},2)="03",(VALUE(LEFT({d},4))-1)&"/12",LEFT({d},4)&"/"&TEXT(VALUE(RIGHT({d},2))-3,"00"))'
    wd.cell(i, ci['Önceki dönem'], '=' + onc(D_)); wd.cell(i, ci['2 önceki dönem'], '=' + onc(P1))
    keys = [f'$B{i}&"|"&{P2}', f'$B{i}&"|"&{P1}', f'$B{i}&"|"&{D_}']
    seri = '&" | "&'.join(f'IFERROR(IF({FK("O", k)}="","—",TEXT({FK("O", k)},"#,##0")),"yok")' for k in keys)
    wd.cell(i, ci['FCF son 3 çeyrek (TTM)'], '=' + seri)
    neg = ','.join(f'IFERROR(N({FK("O", k)})<0,FALSE)' for k in keys)
    ko = ','.join(f'IFERROR({FK("P", k)}="KOŞULLU",FALSE)' for k in keys)
    wd.cell(i, ci['Bölüm F'], f'=IF(AND({neg}),IF(OR({ko}),"ASKIDA (FCF KOŞULLU)","EVET"),"")')
    if r['Puan Durumu'] == 'PUANLANIR':
        wd.cell(i, ci['M1 CFO/NK'], f'=IF(OR({C}="",{NK}="",{NK}=0),"",{C}/{NK})')
        wd.cell(i, ci['M1 Puan'], f'=IF(OR({C}="",{NK}=""),"",IF({NK}<=0,0,IF({C}/{NK}>=1,15,IF({C}/{NK}>=0.5,10,IF({C}/{NK}>0,5,0)))))')
        rng = ','.join(f'{p}{i}' for p in PC)
        wd.cell(i, ci['NULL Metrik #'], '=' + '+'.join(f'IF({p}{i}="",1,0)' for p in PC))
        wd.cell(i, ci['Funnel MIN'], f'=SUM({rng})')
        wd.cell(i, ci['Funnel MAX'], f"={col('Funnel MIN')}{i}+" + '+'.join(f'IF({p}{i}="",{w},0)' for p, (_, w) in zip(PC, PUAN)))
        for t, s in (('Tier MIN', 'Funnel MIN'), ('Tier MAX', 'Funnel MAX')):
            S = f'{col(s)}{i}'; wd.cell(i, ci[t], f'=IF({S}>=70,"ANA LİSTE",IF({S}>=60,"İZLEME",IF({S}>=50,"50–59","ALT")))')
        tmn, tmx = f"{col('Tier MIN')}{i}", f"{col('Tier MAX')}{i}"
        if r['Kapsam'] == 'STANDART':
            wd.cell(i, ci['Kategori'], f'=IF({tmn}<>{tmx},"BELİRSİZ",IF({tmn}="50–59","50–59 · BİLANÇO PUANI GEREKLİ",'
                                        f'IF(AND({tmn}="ANA LİSTE",{col("Bölüm F")}{i}="EVET"),"KOŞULLU ANA LİSTE",{tmn})))')
        else:
            wd.cell(i, ci['Kategori'], v('Son Kategori'))
    else:
        for c in ('M1 CFO/NK', 'M1 Puan', 'NULL Metrik #', 'Funnel MIN', 'Funnel MAX', 'Tier MIN', 'Tier MAX'): wd.cell(i, ci[c], v(c))
        wd.cell(i, ci['Kategori'], v('Son Kategori'))
    e1, e2 = f"{col('Eski Funnel MIN')}{i}", f"{col('Eski Funnel MAX')}{i}"
    wd.cell(i, ci['Değişiklik'], f'=IF({col("Kategori")}{i}<>{col("Eski Kategori")}{i},"KATEGORİ",'
                                 f'IF(OR(N({col("Funnel MIN")}{i})<>N({e1}),N({col("Funnel MAX")}{i})<>N({e2})),"PUAN",""))')
    for j in range(1, len(DC) + 1): wd.cell(i, j).font = F_
    for c in ('CFO TTM (mn TL)', 'Net kâr TTM (mn TL)', 'Eski CFO (Evo, mn TL)'): wd.cell(i, ci[c]).number_format = NUM
    wd.cell(i, ci['CFO TTM (mn TL)']).font = FBL if kay.startswith('Evo') else FG
    for m in MET: wd.cell(i, ci[m]).number_format = RAT
nD = len(FD) + 1
wd.freeze_panes = 'E2'; wd.auto_filter.ref = f'A1:{L(len(DC))}{nD}'
for c in ('Anahtar', 'Önceki dönem', '2 önceki dönem'): wd.column_dimensions[col(c)].hidden = True
for j, w in enumerate([14, 8, 30, 9, 9, 12, 7, 11, 12, 22, 12, 12] + [9, 6] * 11 + [7, 8, 8, 10, 10, 22, 8, 8, 26, 14, 7, 8, 8, 22, 11, 40], 1):
    wd.column_dimensions[L(j)].width = w
wd.row_dimensions[1].height = 54
wd.cell(1, ci['M1 Puan']).comment = Comment('M1 = CFO/Net kâr: Net kâr ≤ 0 → 0; ≥1 → 15; ≥0,5 → 10; >0 → 5. CFO yok → NULL.', 'GEM v6.2')
wd.cell(1, ci['Funnel MAX']).comment = Comment('MIN = puanlar toplamı; MAX = MIN + NULL metriklerin tam puanı (Bölüm E).', 'GEM v6.2')
wd.cell(1, ci['Bölüm F']).comment = Comment('Son üç çeyreğin TTM FCF\'i negatif (FCF sayfasından). Ana Liste ise KOŞULLU ANA LİSTE; '
                                            'üç dönemden biri FCF KOŞULLU ise askıda.', 'GEM v6.2')

GC = ['Kod', 'Dönem', 'Dosya etiketi'] + [c for c in FG_.columns if c not in ('Kod', 'Dönem', 'Dosya etiketi')]
sayfa('FUNNEL_GIRDI', GC, FG_.sort_values(['Kod', 'Dönem'])[GC].values.tolist(), [8, 9, 9] + [14] * (len(GC) - 3),
      bicim={j: TL for j in range(4, len(GC) + 1)}, dondur='C2')

# =============== KODLAR ===============
KOD_L = [
    ('Statü', 'KESİN', 'Tüm kontroller temiz; rakam kullanılır.'), ('Statü', 'TÜRETİLMİŞ', 'Bir bileşen V7 köprü yöntemiyle üretildi; kullanılır, etiketlenir.'),
    ('Statü', 'KOŞULLU', 'Kontrol 9/10/11/13 tetiklendi ve çözülmedi; FCF\'e dayalı yorum yapılmaz.'), ('Statü', 'NULL', 'Rakam yok; NULL kodu nedeni verir. Asla 0 sayılmaz.'),
    ('Kaynak', 'EVO', 'V7 (Evo) rakamı; seri kontrolü temiz, belgeyle tek tek karşılaştırılmadı.'),
    ('Kaynak', 'EVO (belgeyle teyitli)', 'Belge aynı rakamı verdi.'), ('Kaynak', 'BELGE', 'Rakam şirket raporundan (BELGE_TTM / RAPOR_KALEM).'),
    ('Kaynak', 'BELGE (izahname)', 'Halka arz izahnamesindeki finansal tablolardan.'),
    ('Kaynak', 'KARMA', 'Bileşenler farklı kaynaklardan (ör. temettü Evo, diğerleri belge); parantez içinde yazılı.'),
    ('Evo–belge (14g)', 'ISARET_ANOMALISI_COZULEMEDI (K6)', 'Pozitif alım/kira satırı dışlandı, belge çözemedi ve maddi (> |FCF|×%10) → KOŞULLU.'),
    ('NULL kodu', 'NULL_GECMIS_YOK', 'TTM için gerekli önceki yıl raporu yok (yeni halka arz vb.).'),
    ('NULL kodu', 'NULL_BELGE_YOK', 'Rapor bulunamadı (mali yılı takvim dışı şirketler dahil; Açıklama\'da).'),
    ('NULL kodu', 'NULL_OKUNAMADI', 'Rapor var, bileşen güvenle okunamadı.'), ('NULL kodu', 'NULL_CATISMA', 'Evo ≠ belge, açıklanamadı (Funnel girdileri).'),
    ('Evo–belge (14g)', 'BIREBIR', 'Belge ile Evo katsayısız aynı.'), ('Evo–belge (14g)', 'TMS29_KATSAYI_ACIKLIYOR', 'Belge × TMS 29 katsayısı = Evo.'),
    ('Evo–belge (14g)', 'YENI_KARSILASTIRMALI_TUTTU', 'Daha yeni raporun karşılaştırmalı sütunu Evo\'yu tutuyor.'),
    ('Evo–belge (14g)', 'EVO_YENIDEN_BAZLAMAMIS', 'Evo rakamı belgenin çevrilmemiş hâli; belge çevrilerek kullanıldı.'),
    ('Evo–belge (14g)', 'EVO_ESLEME_HATASI', 'Evo yanlış satır/işaret/netleştirme (Kontrol 14 (d)); belge esas.'),
    ('Evo–belge (14g)', 'YENIDEN_DUZENLEME (K11)', 'Şirket geçmişi yeniden düzenlemiş; TTM tek esasla kuruldu.'),
    ('Evo–belge (14g)', 'YAKIN', 'Küçük fark; belge rakamı yazıldı.'), ('Evo–belge (14g)', 'ACIKLANAMAYAN_FARK', 'Belge de çözemedi → KOŞULLU, iki değer.'),
    ('Evo–belge (14g)', 'BELGE_YOK', 'Şüpheli satırın raporu bulunamadı → KOŞULLU.'),
    ('Funnel kategori', 'ANA LİSTE / İZLEME / 50–59 / ALT', '≥70 / 60–69 / 50–59 (bilanço puanı gerekli) / <50; MIN ve MAX farklı tier → BELİRSİZ.'),
    ('Funnel kategori', 'KOŞULLU ANA LİSTE', 'Ana Liste + son üç çeyreğin TTM FCF\'i negatif (Bölüm F).'),
]
sayfa('KODLAR', ['Alan', 'Kod', 'Anlamı'], KOD_L, [18, 30, 110], dondur='A2')

# =============== ANA ===============
unv = FD.drop_duplicates('Kod', keep='last').set_index('Kod')['Ünvan']
kaps = FD.sort_values('Dönem').drop_duplicates('Kod', keep='last').set_index('Kod')['Kapsam']
tip = N.drop_duplicates('Kod', keep='last').set_index('Kod')[['Tip', 'Sektör']]
KOD = sorted(set(N.Kod) | set(FD.Kod))
wa = wb.create_sheet('ANA', 0)
wa['A1'] = 'BIST — FCF TTM ve Funnel V8.2 serileri (MASTER · GEM v6.2)'; wa['A1'].font = FT
wa['A2'] = ('FCF: mn TL, Haziran 2026 satın alma gücü · Funnel: MIN puanı (BELİRSİZ ise MAX FUNNEL sayfasında) · '
            'Hücreler formülle FCF ve FUNNEL sayfalarından gelir; yeni dönem için sütun kopyalanır (OKUBENI).'); wa['A2'].font = FI
grp = [('Şirket', 5), ('FCF · TTM (mn TL)', 6), ('FCF statü', 6), ('Funnel MIN', 5), ('Funnel kategori', 5)]
c0 = 1
for g, n in grp:
    wa.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + n - 1)
    x = wa.cell(3, c0, g); x.font = FB; x.fill = GRP; x.alignment = Alignment(horizontal='center'); c0 += n
AC = ['Kod', 'Ünvan', 'Kapsam', 'Tip', 'Sektör'] + DON + DON + FDON + FDON
for j, c in enumerate(AC, 1):
    x = wa.cell(4, j, c); x.font = FH; x.fill = HDR; x.alignment = Alignment(wrap_text=True, horizontal='center')
BLOK = [('FCF', 'O', nF, 6), ('FCF', 'P', nF, 6), ('FUNNEL', col('Funnel MIN'), nD, 5), ('FUNNEL', col('Kategori'), nD, 5)]
for i, kod in enumerate(KOD, 5):
    wa.cell(i, 1, kod); wa.cell(i, 2, temiz(unv.get(kod))); wa.cell(i, 3, kaps.get(kod, 'FCF kapsamı'))
    if kod in tip.index: wa.cell(i, 4, temiz(tip.loc[kod, 'Tip'])); wa.cell(i, 5, temiz(tip.loc[kod, 'Sektör']))
    j = 6
    for sh, c, n, k in BLOK:
        for _ in range(k):
            x = f'INDEX({sh}!${c}$2:${c}${n},MATCH($A{i}&"|"&{L(j)}$4,{sh}!$A$2:$A${n},0))'
            wa.cell(i, j, f'=IFERROR(IF({x}="","",{x}),"")')
            if sh == 'FCF' and c == 'O': wa.cell(i, j).number_format = NUM
            j += 1
    for jj in range(1, len(AC) + 1): wa.cell(i, jj).font = F_
nA = len(KOD) + 4
wa.freeze_panes = 'C5'; wa.auto_filter.ref = f'A4:{L(len(AC))}{nA}'
for j, w in enumerate([8, 34, 12, 10, 22] + [11] * 6 + [11] * 6 + [8] * 5 + [15] * 5, 1): wa.column_dimensions[L(j)].width = w
for t, clr in (('KOŞULLU', 'FFEB9C'), ('NULL', 'D9D9D9')):
    wa.conditional_formatting.add(f'L5:Q{nA}', CellIsRule(operator='equal', formula=[f'"{t}"'], fill=PatternFill('solid', fgColor=clr)))
wa.conditional_formatting.add(f'R5:V{nA}', CellIsRule(operator='greaterThanOrEqual', formula=['70'], fill=PatternFill('solid', fgColor='C6EFCE')))
for t, clr in (('ANA LİSTE', 'C6EFCE'), ('KOŞULLU ANA LİSTE', 'FFEB9C'), ('İZLEME', 'DDEBF7'), ('BELİRSİZ', 'FCE4D6')):
    wa.conditional_formatting.add(f'W5:AA{nA}', CellIsRule(operator='equal', formula=[f'"{t}"'], fill=PatternFill('solid', fgColor=clr)))

# =============== OKUBENI ===============
OK = [
    ('Bu dosya', 'BIST FCF ve Funnel için TEK KAYNAK (master). Kural seti: GEM v6.2 (FCF v2.2, Funnel V8.2). Dönemler: FCF 2025/03–2026/06, '
                 'Funnel 2025/06–2026/06. Kilitli sürüm; yeni dönem eklenerek büyür.'),
    ('ANA', 'Şirket başına bir satır: FCF TTM + statü serisi, Funnel MIN + kategori serisi. "Son 3 dönem pozitif" gibi sorgular buradan.'),
    ('FCF', 'Şirket × dönem. FCF_STD = CFO_TTM − CAPEX_STD − |Kira anapara|; CAPEX_STD = |MDV+MODV alımı| + |YAGM alımı|; '
            'holdingde FCF_HLD = FCF_STD + alınan temettü (yatırım bölümü). Mavi = girdi, siyah = formül. Statü, Kaynak, NULL kodu, '
            'Evo–belge kodu (Kontrol 14 g) ve Açıklama her satırda.'),
    ('BELGE_TTM', 'Belgeden kurulan her bileşenin TTM parçaları: A (cari YTD) + B (önceki yıl sonu) − C (önceki yıl aynı dönem), '
                  'katsayılar ve kanıt (KAP bildirim no, sayfa, satır). "FCF\'te kullanıldı = EVET" satırları FCF sayfasının girdisidir.'),
    ('RAPOR_KALEM', 'Belgeler: okunan her rapor kalemi, raporda yazdığı gibi (TL, çevrilmemiş), kaynağı ve bağlantısıyla '
                    '(KAP XBRL/PDF, şirket sitesi, Evo belge havuzu, izahname).'),
    ('ELLE_KARAR', 'Belge elle okunarak verilen kararlar ve gerekçeleri (XBRL hatası, PDF satırı seçimi vb.).'),
    ('NOTLAR', 'Şirket bazında çözülemeyen satırların gerekçesi, belge kararları, kod değişiklikleri ve genel bulgular.'),
    ('TMS29', 'Belge rakamlarını Haziran 2026 satın alma gücüne çeviren katsayılar ve dayanağı.'),
    ('FUNNEL', 'Funnel V8.2: M1–M12 değerleri ve puanları, NULL sayısı, MIN/MAX, tier, kategori. M1\'in CFO\'su FCF sayfasındaki '
               'CFO\'dan farklıysa o kullanılır (yeşil). Bölüm F (Ana Liste + son üç çeyrek TTM FCF negatif → KOŞULLU ANA LİSTE) '
               'FCF sayfasından formülle. Eski tarama değerleri yan sütunlarda.'),
    ('FUNNEL_GIRDI', 'Funnel\'ın Evo ham girdileri (TL).'),
    ('KODLAR', 'Statü, kaynak, NULL kodu, Evo–belge kodu ve Funnel kategori listeleri.'),
    ('YENİ DÖNEM EKLEME', '1) Yeni dönemin tam taraması/hesabı çalıştırılır (Evo + Kontrol 13/14 + belge). '
                          '2) FCF sayfasının altına yeni satırlar eklenir (Anahtar = KOD|YYYY/MM; formül sütunları üst satırdan kopyalanır). '
                          '3) Belgeden okunanlar BELGE_TTM ve RAPOR_KALEM\'e, elle kararlar ELLE_KARAR\'a, notlar NOTLAR\'a eklenir. '
                          '4) Funnel sonuçları FUNNEL ve FUNNEL_GIRDI\'ye eklenir. '
                          '5) ANA\'da FCF ve Funnel bloklarına birer sütun eklenir; başlığa yeni dönem yazılır, formül soldaki sütundan kopyalanır. '
                          '6) Formül aralıkları (…$2:$X$n) yeni satırları kapsayacak şekilde genişletilir. Eski satırlar değiştirilmez; '
                          'düzeltme gerekirse NOTLAR\'a kayıt düşülür.'),
    ('Kapsam', 'Bankalar, sigorta, GYO ve benzeri (Tier 0, Tier 0,5, standart dışı) FCF kapsamı dışında; Funnel satırları var. '
               'Kodu değişen 5 eski kod (KOZAL, KOZAA, IPEKE, MARKA, SNKRN) çıkarıldı (NOTLAR).'),
]
wo = wb.create_sheet('OKUBENI', 1)
for j, c in enumerate(['Konu', 'Açıklama'], 1):
    x = wo.cell(1, j, c); x.font = FH; x.fill = HDR
for i, (a, b) in enumerate(OK, 2):
    wo.cell(i, 1, a).font = FB; x = wo.cell(i, 2, b); x.font = F_; x.alignment = Alignment(wrap_text=True, vertical='top')
wo.column_dimensions['A'].width = 20; wo.column_dimensions['B'].width = 130

wb.save('FCF_MASTER.xlsx')
print('yazıldı:', len(KOD), 'şirket ·', nF - 1, 'FCF ·', len(bt), 'BELGE_TTM ·', len(rk), 'RAPOR_KALEM ·', len(er), 'ELLE ·', nD - 1, 'FUNNEL')
print(N['NULL kodu'].value_counts().to_dict()); print(N['Evo–belge kodu (14g)'].fillna('').str.split(':').str[-1].value_counts().head(12).to_dict())
