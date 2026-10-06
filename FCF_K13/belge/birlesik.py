"""Tek dosya: nihai FCF (bileşenleriyle, formüllü) + 5 dönemlik Funnel V8.2 (bileşenleriyle, M1 güncel CFO ile, formüllü)
+ ana sayfada şirket bazında Funnel ve FCF serileri.
Kullanım (FCF_K13 içinden): python3 belge/birlesik.py <GEM_20xxQx_Funnel_V8_2_FCF_V2_1.xlsx ...>
Çıktı: FCF_FUNNEL_birlesik.xlsx"""
import sys, re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.comments import Comment

DON = ['2025/03', '2025/06', '2025/09', '2025/12', '2026/03', '2026/06']
FDON = ['2025/06', '2025/09', '2025/12', '2026/03', '2026/06']
Q = {'Q1': '03', 'Q2': '06', 'Q3': '09', 'Q4': '12'}
PUAN = [('M1', 15), ('M2', 10), ('M3', 5), ('M5', 12), ('M6', 11), ('M7', 10), ('M8', 10), ('M9', 8), ('M10', 7), ('M11', 7),
        ('M12', 5)]

F_ = Font(name='Arial', size=10); FB = Font(name='Arial', size=10, bold=True); FH = Font(name='Arial', size=10, bold=True, color='FFFFFF')
FT = Font(name='Arial', size=14, bold=True); FG = Font(name='Arial', size=10, color='008000'); FBL = Font(name='Arial', size=10, color='0000FF')
HDR = PatternFill('solid', fgColor='1F3864'); GRP = PatternFill('solid', fgColor='D9E1F2')
NUM = '#,##0.0;[Red]-#,##0.0;-'; INT = '0;[Red]-0;0'; RAT = '0.00;[Red]-0.00'

# ---------------- veri ----------------
N = pd.read_excel('FCF_TTM_nihai_belgeli.xlsx', sheet_name='FCF_TTM').sort_values(['Kod', 'Dönem'], ignore_index=True)
assert (N.groupby('Kod').size() == 6).all()  # her şirket 6 dönem, bloklar ardışık
# CAPEX ayrımı: V7 Bileşenler (formülde kullanılan) + belge satırları (null/kosullu csv)
B = pd.read_excel('girdi/FCF_TTM_2025-03_2026-06_v7_nihai.xlsx', sheet_name='Bileşenler')
B = B[B['Bileşen'].isin(['MDV+MODV alımı', 'YAGM alımı'])].pivot_table(index=['Kod', 'Dönem'], columns='Bileşen',
                                                                        values='Formülde (mn TL)', aggfunc='first')
belge = {}
for f in ('belge/null_ttm.csv', 'belge/kosullu_sonuc.csv'):
    for _, r in pd.read_csv(f, keep_default_na=False).iterrows():
        def u(x):
            if x == '': return None
            x = float(x); return -x if x < 0 else 0.0  # K6: pozitif TTM alım satırı kullanılmaz
        belge[(r['Kod'], r['Dönem'])] = (u(r['MDV+MODV TTM']), u(r['YAGM TTM']))
mdv, yagm = [], []
for _, r in N.iterrows():
    k = (r['Kod'], r['Dönem']); cap = r['CAPEX_STD']
    if pd.isna(cap): mdv.append(None); yagm.append(None); continue
    cand = []
    if str(r['Kaynak']).startswith('BELGE') or 'teyitli' in str(r['Kaynak']):
        if k in belge and None not in belge[k]: cand.append(belge[k])
    if k in B.index:
        a, b = B.loc[k].get('MDV+MODV alımı'), B.loc[k].get('YAGM alımı')
        cand.append((0.0 if pd.isna(a) else abs(a), 0.0 if pd.isna(b) else abs(b)))
    ok = next((c for c in cand if abs(c[0] + c[1] - cap) < 0.01), None)
    if ok is None:  # ayrım bulunamazsa YAGM bilinen değer, MDV = kalan
        y = cand[0][1] if cand else 0.0
        ok = (cap - y, y)
    mdv.append(ok[0]); yagm.append(ok[1])
N['MDV+MODV'] = mdv; N['YAGM'] = yagm

# Funnel dosyaları
FD, FG_ = [], []
for f in sys.argv[1:]:
    y, q = re.search(r'GEM_(\d{4})(Q\d)', f).groups(); et = f'{y}/{Q[q]}'
    son = pd.Timestamp(f'{y}-{Q[q]}-01') + pd.offsets.MonthEnd(0)
    d = pd.read_excel(f, sheet_name='Funnel_Detay', header=2); d = d[d.Kod.notna()]
    g = pd.read_excel(f, sheet_name='Evo_Girdileri', header=2).rename(columns={'hisse_senedi_kodu': 'Kod'})
    g = g[g.Kod.notna()]
    def takvim(yay):
        yay = pd.Timestamp(yay)
        if pd.isna(yay) or (yay - son).days <= 110: return et
        x = yay - pd.Timedelta(days=20)
        qe = pd.Timestamp(x.year, ((x.month - 1) // 3) * 3 + 1, 1) - pd.Timedelta(days=1)
        return f'{qe.year}/{qe.month:02d}'
    g['Dönem'] = g['yayinlanma_tarihi_utc'].map(takvim); g['Dosya etiketi'] = et
    d = d.merge(g[['Kod', 'Dönem', 'Dosya etiketi', 'cfo_ttm', 'net_total_cur', 'yayinlanma_tarihi_utc']], on='Kod', how='left')
    d['Dönem'] = d['Dönem'].fillna(et); d['Dosya etiketi'] = d['Dosya etiketi'].fillna(et)
    FD.append(d); FG_.append(g)
FD = pd.concat(FD, ignore_index=True).drop_duplicates(['Kod', 'Dönem'], keep='last')
FG_ = pd.concat(FG_, ignore_index=True).drop_duplicates(['Kod', 'Dönem'], keep='last')
NC = N.set_index(['Kod', 'Dönem'])

# şirket listesi
unv = pd.concat([FD[['Kod', 'Ünvan']]]).drop_duplicates('Kod', keep='last').set_index('Kod')['Ünvan']
kaps = FD.sort_values('Dönem').drop_duplicates('Kod', keep='last').set_index('Kod')['Kapsam']
tip = N.drop_duplicates('Kod', keep='last').set_index('Kod')[['Tip', 'Sektör']]
KOD = sorted(set(N.Kod) | set(FD.Kod))

wb = Workbook()

def baslik(ws, row, cols, fill=HDR, font=FH):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row, j, c); x.font = font; x.fill = fill; x.alignment = Alignment(wrap_text=True, vertical='center')

def genislik(ws, w):
    for j, v in enumerate(w, 1): ws.column_dimensions[L(j)].width = v

# ---------------- FCF_DETAY ----------------
wf = wb.active; wf.title = 'FCF_DETAY'
FC = ['Anahtar', 'Kod', 'Tip', 'Sektör', 'Dönem', 'CFO_TTM', '|MDV+MODV alımı|', '|YAGM alımı|', 'CAPEX_STD', '|Kira anapara|',
      'Alınan temettü (CFO dışı)', 'FCF_STD (①)', 'FCF_HLD (②)', 'FCF türü', 'FCF · TTM (mn TL)', 'Statü', 'Kaynak',
      'Piyasa değeri (mn TL)', 'P/FCF', 'Açıklama', 'V7 FCF', 'V7 statü', "V7'ye göre değişti"]
baslik(wf, 1, FC)
for i, (_, r) in enumerate(N.iterrows(), 2):
    v = lambda c: None if pd.isna(r[c]) else r[c]
    wf.cell(i, 1, f"{r['Kod']}|{r['Dönem']}"); wf.cell(i, 2, r['Kod']); wf.cell(i, 3, v('Tip')); wf.cell(i, 4, v('Sektör'))
    wf.cell(i, 5, r['Dönem']); wf.cell(i, 6, v('CFO_TTM')); wf.cell(i, 7, v('MDV+MODV')); wf.cell(i, 8, v('YAGM'))
    wf.cell(i, 9, f'=IF(OR(F{i}="",G{i}="",H{i}=""),"",G{i}+H{i})'); wf.cell(i, 10, v('|Kira anapara|'))
    wf.cell(i, 11, v('Alınan temettü (CFO dışı)'))
    wf.cell(i, 12, f'=IF(OR(F{i}="",I{i}="",J{i}=""),"",F{i}-I{i}-J{i})')
    wf.cell(i, 13, f'=IF(OR(L{i}="",N{i}<>"FCF_HLD"),"",L{i}+N(K{i}))')
    wf.cell(i, 14, v('FCF türü'))
    wf.cell(i, 15, f'=IF(L{i}="","",IF(N{i}="FCF_HLD",M{i},L{i}))')
    wf.cell(i, 16, v('Statü')); wf.cell(i, 17, v('Kaynak')); wf.cell(i, 18, v('Piyasa değeri (mn TL)'))
    wf.cell(i, 19, f'=IF(OR(O{i}="",R{i}=""),"",IF(O{i}>0,R{i}/O{i},""))')
    wf.cell(i, 20, v('Açıklama')); wf.cell(i, 21, v('V7 FCF')); wf.cell(i, 22, v('V7 statü')); wf.cell(i, 23, v("V7'ye göre değişti"))
    for j in range(1, 24):
        c = wf.cell(i, j); c.font = F_
        if j in (6, 7, 8, 10, 11, 18, 21): c.font = FBL
        if j in (6, 7, 8, 9, 10, 11, 12, 13, 15, 18, 21): c.number_format = NUM
        if j == 19: c.number_format = '0.0x'
nF = len(N) + 1
wf.freeze_panes = 'F2'; wf.auto_filter.ref = f'A1:{L(len(FC))}{nF}'; wf.column_dimensions['A'].hidden = True
genislik(wf, [14, 8, 10, 22, 9] + [12] * 8 + [10, 13, 12, 22, 12, 8, 60, 11, 11, 14])
wf.row_dimensions[1].height = 42

# ---------------- FUNNEL_DETAY ----------------
wd = wb.create_sheet('FUNNEL_DETAY')
MET = ['M1 CFO/NK', 'M2 Net Borç/FAVÖK', 'M3 Cari Oran', 'M5 ROE', 'M6 Brüt Marj', 'M7 Satış Büyümesi', 'M8 Özkaynak Büyümesi',
       'M9 OPEX/Brüt Kâr', 'M10 R', 'M11 FAVÖK/NK', 'M12 Marj Trendi']
DC = (['Anahtar', 'Kod', 'Ünvan', 'Dönem', 'Dosya etiketi', 'Kapsam', 'Kritik Eksik #', 'Puan Durumu',
       'CFO TTM (güncel, mn TL)', 'CFO kaynağı', 'Net kâr TTM (mn TL)', 'Eski CFO (Evo, mn TL)'] +
      sum([[m, m.split()[0] + ' Puan'] for m in MET], []) +
      ['NULL Metrik #', 'Funnel MIN', 'Funnel MAX', 'Tier MIN', 'Tier MAX', 'Kategori (güncel)', 'FCF son 3 çeyrek (TTM)',
       'Bölüm F', 'Eski M1 Puan', 'Eski Funnel MIN', 'Eski Funnel MAX', 'Eski Kategori', 'Değişiklik', 'Not', 'FCF satırı'])
baslik(wd, 1, DC)
ci = {c: j for j, c in enumerate(DC, 1)}
col = lambda c: L(ci[c])
PC = [col(m.split()[0] + ' Puan') for m in MET]
FD = FD.sort_values(['Kod', 'Dönem'])
for i, (_, r) in enumerate(FD.iterrows(), 2):
    v = lambda c: None if c not in r or pd.isna(r[c]) else r[c]
    kod, don = r['Kod'], r['Dönem']
    ecfo = None if pd.isna(r['cfo_ttm']) else r['cfo_ttm'] / 1e6
    ncfo, kay = ecfo, 'Evo (eski tarama)'
    if (kod, don) in NC.index and not pd.isna(NC.loc[(kod, don), 'CFO_TTM']):
        x = NC.loc[(kod, don)]
        if ecfo is None or abs(x['CFO_TTM'] - ecfo) > max(0.005 * abs(x['CFO_TTM']), 0.5):
            ncfo, kay = x['CFO_TTM'], 'Nihai FCF dosyası (' + str(x['Kaynak']) + ')'
    vals = {'Anahtar': f'{kod}|{don}', 'Kod': kod, 'Ünvan': v('Ünvan'), 'Dönem': don, 'Dosya etiketi': v('Dosya etiketi'),
            'Kapsam': v('Kapsam'), 'Kritik Eksik #': v('Kritik Eksik #'), 'Puan Durumu': v('Puan Durumu'),
            'CFO TTM (güncel, mn TL)': ncfo, 'CFO kaynağı': kay,
            'Net kâr TTM (mn TL)': None if pd.isna(r['net_total_cur']) else r['net_total_cur'] / 1e6,
            'Eski CFO (Evo, mn TL)': ecfo, 'Eski M1 Puan': v('M1 Puan'), 'Eski Funnel MIN': v('Funnel MIN'),
            'Eski Funnel MAX': v('Funnel MAX'), 'Eski Kategori': v('Son Kategori'), 'Not': v('Not')}
    for m in MET[1:]:
        vals[m] = v(m); vals[m.split()[0] + ' Puan'] = v(m.split()[0] + ' Puan')
    for c, x in vals.items(): wd.cell(i, ci[c], x)
    puanlanir = r['Puan Durumu'] == 'PUANLANIR'
    C, NK = f"{col('CFO TTM (güncel, mn TL)')}{i}", f"{col('Net kâr TTM (mn TL)')}{i}"
    if puanlanir:
        wd.cell(i, ci['M1 CFO/NK'], f'=IF(OR({C}="",{NK}="",{NK}=0),"",{C}/{NK})')
        wd.cell(i, ci['M1 Puan'], f'=IF(OR({C}="",{NK}=""),"",IF({NK}<=0,0,IF({C}/{NK}>=1,15,IF({C}/{NK}>=0.5,10,IF({C}/{NK}>0,5,0)))))')
        rng = ','.join(f'{p}{i}' for p in PC)
        wd.cell(i, ci['NULL Metrik #'], '=' + '+'.join(f'IF({p}{i}="",1,0)' for p in PC))
        wd.cell(i, ci['Funnel MIN'], f'=SUM({rng})')
        wd.cell(i, ci['Funnel MAX'], f"={col('Funnel MIN')}{i}+" + '+'.join(f'IF({p}{i}="",{w},0)' for p, (_, w) in zip(PC, PUAN)))
        for t, s in (('Tier MIN', 'Funnel MIN'), ('Tier MAX', 'Funnel MAX')):
            S = f'{col(s)}{i}'
            wd.cell(i, ci[t], f'=IF({S}>=70,"ANA LİSTE",IF({S}>=60,"İZLEME",IF({S}>=50,"50–59","ALT")))')
        tmn, tmx = f"{col('Tier MIN')}{i}", f"{col('Tier MAX')}{i}"
        kat = (f'=IF({tmn}<>{tmx},"BELİRSİZ",IF({tmn}="50–59","50–59 · BİLANÇO PUANI GEREKLİ",'
               f'IF(AND({tmn}="ANA LİSTE",{col("Bölüm F")}{i}="EVET"),"KOŞULLU ANA LİSTE",{tmn})))')
        wd.cell(i, ci['Kategori (güncel)'], kat if r['Kapsam'] == 'STANDART' else v('Son Kategori'))
        if r['Kapsam'] != 'STANDART':
            wd.cell(i, ci['Kategori (güncel)'], v('Son Kategori'))
    else:
        for c in ('M1 CFO/NK', 'M1 Puan', 'NULL Metrik #', 'Funnel MIN', 'Funnel MAX', 'Tier MIN', 'Tier MAX'):
            wd.cell(i, ci[c], v(c))
        wd.cell(i, ci['Kategori (güncel)'], v('Son Kategori'))
    # Bölüm F: son üç çeyreğin TTM FCF'i (FCF_DETAY'dan) ve statüleri
    R = f"{col('FCF satırı')}{i}"
    wd.cell(i, ci['FCF satırı'], f'=IFERROR(MATCH(A{i},FCF_DETAY!$A$2:$A${nF},0),"")')
    if don in DON and DON.index(don) >= 2:
        O = lambda k: f'INDEX(FCF_DETAY!$O$2:$O${nF},{R}-{k})'
        P = lambda k: f'INDEX(FCF_DETAY!$P$2:$P${nF},{R}-{k})'
        seri = '&" | "&'.join(f'IF({O(k)}="","—",TEXT({O(k)},"#,##0"))' for k in (2, 1, 0))
        wd.cell(i, ci['FCF son 3 çeyrek (TTM)'], f'=IF({R}="","",{seri})')
        neg = ','.join(f'N({O(k)})<0' for k in (2, 1, 0))
        kos = ','.join(f'{P(k)}="KOŞULLU"' for k in (2, 1, 0))
        wd.cell(i, ci['Bölüm F'], f'=IF({R}="","",IF(AND({neg}),IF(OR({kos}),"ASKIDA (FCF KOŞULLU)","EVET"),""))')
    e1, e2 = f"{col('Eski Funnel MIN')}{i}", f"{col('Eski Funnel MAX')}{i}"
    wd.cell(i, ci['Değişiklik'], f'=IF({col("Kategori (güncel)")}{i}<>{col("Eski Kategori")}{i},"KATEGORİ",'
                                 f'IF(OR(N({col("Funnel MIN")}{i})<>N({e1}),N({col("Funnel MAX")}{i})<>N({e2})),"PUAN",""))')
    for j in range(1, len(DC) + 1):
        c = wd.cell(i, j); c.font = F_
    for c in ('CFO TTM (güncel, mn TL)', 'Net kâr TTM (mn TL)', 'Eski CFO (Evo, mn TL)'):
        wd.cell(i, ci[c]).number_format = NUM
    wd.cell(i, ci['CFO TTM (güncel, mn TL)']).font = FBL if kay.startswith('Evo') else FG
    for m in MET: wd.cell(i, ci[m]).number_format = RAT
nD = len(FD) + 1
wd.freeze_panes = 'E2'; wd.auto_filter.ref = f'A1:{L(len(DC))}{nD}'; wd.column_dimensions['A'].hidden = True
wd.column_dimensions[col('FCF satırı')].hidden = True
genislik(wd, [14, 8, 30, 9, 9, 12, 7, 11, 12, 22, 12, 12] + [9, 6] * 11 + [7, 8, 8, 10, 10, 22, 26, 14, 7, 8, 8, 22, 11, 40])
wd.row_dimensions[1].height = 54

# ---------------- FUNNEL_GIRDILER (Evo ham, TL) ----------------
wg = wb.create_sheet('FUNNEL_GIRDILER')
GC = ['Kod', 'Dönem', 'Dosya etiketi'] + [c for c in FG_.columns if c not in ('Kod', 'Dönem', 'Dosya etiketi')]
baslik(wg, 1, GC)
for i, r in enumerate(FG_.sort_values(['Kod', 'Dönem'])[GC].itertuples(index=False), 2):
    for j, x in enumerate(r, 1):
        c = wg.cell(i, j, None if (not isinstance(x, str) and pd.isna(x)) else x); c.font = F_
        if isinstance(x, float): c.number_format = '#,##0;[Red]-#,##0'
wg.freeze_panes = 'C2'; wg.auto_filter.ref = f'A1:{L(len(GC))}{len(FG_) + 1}'

# ---------------- ANA ----------------
wa = wb.create_sheet('ANA', 0)
wa['A1'] = 'BIST — Funnel V8.2 ve FCF TTM serileri (GEM v6.2)'; wa['A1'].font = FT
wa['A2'] = ('Funnel: eski toplu tarama (Evo, 03.09.2026), M1 güncel CFO ile yeniden hesaplandı · FCF: nihai belge doğrulamalı '
            'tablo (mn TL, Haziran 2026 satın alma gücü) · Değerler formülle ayrıntı sayfalarından gelir.')
wa['A2'].font = Font(name='Arial', size=9, italic=True)
grp = [('Şirket', 5), ('Funnel MIN (puan)', 5), ('Kategori', 5), ('FCF · TTM (mn TL)', 6), ('FCF statü', 6), ('Son', 2)]
c0 = 1
for g, n in grp:
    wa.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + n - 1)
    x = wa.cell(3, c0, g); x.font = FB; x.fill = GRP; x.alignment = Alignment(horizontal='center')
    c0 += n
AC = (['Kod', 'Ünvan', 'Kapsam', 'Tip', 'Sektör'] + FDON + FDON + DON + DON + ['P/FCF (son)', 'Değişiklik (son dönem)'])
baslik(wa, 4, AC)
look = lambda sh, c, n, key: f'INDEX({sh}!${c}${2}:${c}${n},MATCH({key},{sh}!$A$2:$A${n},0))'
for i, kod in enumerate(KOD, 5):
    wa.cell(i, 1, kod)
    wa.cell(i, 2, unv.get(kod)); wa.cell(i, 3, kaps.get(kod, 'FCF kapsamı'))
    if kod in tip.index:
        wa.cell(i, 4, tip.loc[kod, 'Tip']); wa.cell(i, 5, tip.loc[kod, 'Sektör'])
    H0 = len(AC) + 1  # gizli yardımcı sütunlar: FCF blok başı + 5 Funnel satırı
    wa.cell(i, H0, f'=IFERROR(MATCH($A{i}&"|2025/03",FCF_DETAY!$A$2:$A${nF},0),"")')
    for k, d in enumerate(FDON):
        wa.cell(i, H0 + 1 + k, f'=IFERROR(MATCH($A{i}&"|{d}",FUNNEL_DETAY!$A$2:$A${nD},0),"")')
    j = 6
    for kk, (c, k0) in enumerate(((col('Funnel MIN'), 0), (col('Kategori (güncel)'), 0))):
        for k, d in enumerate(FDON):
            h = f'{L(H0 + 1 + k)}{i}'; x = f'INDEX(FUNNEL_DETAY!${c}$2:${c}${nD},{h})'
            wa.cell(i, j, f'=IF({h}="","",IF({x}="","",{x}))'); j += 1
    for c in ('O', 'P'):
        for k, d in enumerate(DON):
            h = f'{L(H0)}{i}'; x = f'INDEX(FCF_DETAY!${c}$2:${c}${nF},{h}+{k})'
            wa.cell(i, j, f'=IF({h}="","",IF({x}="","",{x}))')
            if c == 'O': wa.cell(i, j).number_format = NUM
            j += 1
    h = f'{L(H0)}{i}'; x = f'INDEX(FCF_DETAY!$S$2:$S${nF},{h}+5)'
    wa.cell(i, j, f'=IF({h}="","",IF({x}="","",{x}))'); wa.cell(i, j).number_format = '0.0x'
    h = f'{L(H0 + 5)}{i}'; x = f"INDEX(FUNNEL_DETAY!${col('Değişiklik')}$2:${col('Değişiklik')}${nD},{h})"
    wa.cell(i, j + 1, f'=IF({h}="","",IF({x}="","",{x}))')
    for jj in range(1, len(AC) + 1): wa.cell(i, jj).font = F_
nA = len(KOD) + 4
wa.freeze_panes = 'C5'; wa.auto_filter.ref = f'A4:{L(len(AC))}{nA}'
for k in range(6): wa.column_dimensions[L(len(AC) + 1 + k)].hidden = True
genislik(wa, [8, 34, 12, 10, 22] + [8] * 5 + [15] * 5 + [11] * 6 + [11] * 6 + [9, 12])
wa.row_dimensions[4].height = 30
# koşullu biçim: kategori renkleri
kr = f'K5:O{nA}'
for t, clr in (('ANA LİSTE', 'C6EFCE'), ('KOŞULLU ANA LİSTE', 'FFEB9C'), ('İZLEME', 'DDEBF7'), ('BELİRSİZ', 'FCE4D6')):
    wa.conditional_formatting.add(kr, CellIsRule(operator='equal', formula=[f'"{t}"'], fill=PatternFill('solid', fgColor=clr)))
fr = f'F5:J{nA}'
wa.conditional_formatting.add(fr, CellIsRule(operator='greaterThanOrEqual', formula=['70'], fill=PatternFill('solid', fgColor='C6EFCE')))
sr = f'W5:AB{nA}'
for t, clr in (('KOŞULLU', 'FFEB9C'), ('NULL', 'D9D9D9')):
    wa.conditional_formatting.add(sr, CellIsRule(operator='equal', formula=[f'"{t}"'], fill=PatternFill('solid', fgColor=clr)))

# ---------------- OKUBENI ----------------
wo = wb.create_sheet('OKUBENI', 1)
OK = [
    ('Ne', 'Tek dosya: nihai FCF TTM (2025/03–2026/06, bileşenleriyle) + Funnel V8.2 (2025/06–2026/06, metrik ve puanlarıyla).'),
    ('ANA', 'Şirket başına bir satır: Funnel MIN puanı ve kategori serisi (5 dönem), FCF TTM ve statü serisi (6 dönem), son P/FCF. '
            'Hücreler formülle FUNNEL_DETAY ve FCF_DETAY sayfalarından gelir; filtre ve sıralama başlık satırında.'),
    ('FCF_DETAY', 'FCF_STD = CFO − CAPEX_STD − |kira anapara|; CAPEX_STD = |MDV+MODV alımı| + |YAGM alımı|; holdingde FCF_HLD = '
                  'FCF_STD + alınan temettü (yatırım bölümü). Mavi = girdi (Evo ya da belge), siyah = formül. mn TL, Haziran 2026 '
                  'satın alma gücü. Statü / Kaynak / Açıklama: FCF_TTM_nihai_belgeli.xlsx ile aynı.'),
    ('FUNNEL_DETAY', 'Eski toplu taramanın (Evo, çekim 03.09.2026) metrikleri ve puanları. Değişen tek girdi M1\'in CFO\'su: nihai '
                     'FCF dosyasındaki CFO eski Evo CFO\'sundan farklıysa o kullanıldı (yeşil yazı; kaynağı "CFO kaynağı" sütununda). '
                     'M1 puanı, NULL metrik sayısı, MIN/MAX, tier ve kategori formülle yeniden hesaplanır. Eski değerler yanında.'),
    ('Kategori', '≥70 ANA LİSTE · 60–69 İZLEME · 50–59 BİLANÇO PUANI GEREKLİ · <50 ALT · MIN ve MAX farklı tier → BELİRSİZ. '
                 'Bölüm F: ANA LİSTE ve son üç çeyreğin TTM FCF\'i negatif → KOŞULLU ANA LİSTE; bu üç dönemden biri FCF KOŞULLU '
                 'ise kural askıda. Eski tarama Bölüm F\'i uygulamamıştı.'),
    ('Dönem', 'Takvim dönemi. Mali yılı takvim yılı olmayan şirketlerde eski dosya kendi mali çeyreğiyle etiketlemiş; burada yayın '
              'tarihinden takvim dönemine çevrildi (eski etiket "Dosya etiketi" sütununda).'),
    ('Kapsam dışı', 'Funnel sayfalarında Tier 0 / Tier 0,5 / standart dışı şirketler var ama FCF yok (FCF çalışması kapsamı dışı). '
                    'FCF\'te 2025/03 var, Funnel\'da yok (eski taramalar 2025/06\'dan başlıyor).'),
    ('Çeyreklik FCF', 'Tek çeyrek FCF bu dosyada yok; yalnız TTM.'),
    ('FUNNEL_GIRDILER', 'Eski taramanın Evo ham girdileri (TL), dönem bazında.'),
]
baslik(wo, 1, ['Konu', 'Açıklama'])
for i, (a, b) in enumerate(OK, 2):
    wo.cell(i, 1, a).font = FB; x = wo.cell(i, 2, b); x.font = F_; x.alignment = Alignment(wrap_text=True, vertical='top')
genislik(wo, [16, 120])
wb.save('FCF_FUNNEL_birlesik.xlsx')
print('yazıldı', len(KOD), 'şirket', nF - 1, 'FCF satırı', nD - 1, 'Funnel satırı')
